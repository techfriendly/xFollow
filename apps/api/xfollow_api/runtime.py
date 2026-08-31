from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
import zipfile
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

import fitz
import boto3
from docx import Document
from psycopg2.extras import Json
from psycopg2.errors import UniqueViolation

from . import db, integrations
from .config import load_runtime_config, repo_root, sanitized_config
from .models import (
    AlertUpdate,
    AlertRuleCreate,
    ComplianceCheckRequest,
    ContractCommentCreate,
    ContractCreate,
    ContractChangeCreate,
    ContractChangeUpdate,
    ContractRiskUpdate,
    ContractUpdate,
    DataProtectionAssessmentUpdate,
    EvidenceCreate,
    FeedbackCreate,
    GeneratedDocumentRequest,
    ImportPcspRequest,
    ImportXreviewRequest,
    ImportXtenderRequest,
    InvoiceCreate,
    InvoiceUpdate,
    ObligationCreate,
    ObligationUpdate,
    PenaltyCalculationRequest,
    ProceedingCreate,
    ProceedingUpdate,
    ResolutionRequest,
    TaskCreate,
    TaskUpdate,
    UserContext,
)
from .security import ROLE_PERMISSIONS

_MEMORY: dict[str, dict[str, dict[str, Any]]] = {
    "contracts": {},
    "obligations": {},
    "alerts": {},
    "evidence": {},
    "checks": {},
    "risks": {},
    "documents": {},
    "penalties": {},
    "invoices": {},
    "proceedings": {},
    "changes": {},
    "feedback": {},
    "data_protection": {},
    "tasks": {},
    "comments": {},
    "alert_rules": {},
    "audit": {},
    "evidence_files": {},
    "xtender": {
        "EXP-DEMO-XTENDER": {
            "id": "EXP-DEMO-XTENDER",
            "tenant_id": "tenant-demo",
            "title": "Servicio de mantenimiento de edificios municipales",
            "unit": "Servicios Generales",
            "data": {
                "file_number": "2026/SERV/001",
                "contracting_body": "Ayuntamiento demo",
                "contract_manager": "Responsable del contrato",
                "cpv_codes": ["50700000-2"],
                "contract_type": "Servicios",
                "procedure": "Abierto",
                "budget": 180000,
                "estimated_value": 216000,
                "start_date": "2026-08-01",
                "end_date": "2027-07-31",
                "extensions": "Una prorroga de 12 meses con preaviso de 2 meses.",
                "special_execution_conditions": [
                    {"title": "Plan de mantenimiento preventivo", "description": "Entrega mensual de partes y plan preventivo."}
                ],
            },
        }
    },
    "pcsp": {
        "PCSP-DEMO-ADJ": {
            "id": "PCSP-DEMO-ADJ",
            "tenant_id": "tenant-demo",
            "expediente": "2026/ADJ/001",
            "title": "Suministro de equipamiento informatico adjudicado",
            "status": "ADJ",
            "cpv": "30200000-1",
            "contracting_body": "Entidad publica demo",
            "budget_without_tax": 95000,
            "budget_with_tax": 114950,
            "currency": "EUR",
        }
    },
    "xreview": {
        "HOF-DEMO-XREVIEW": {
            "id": "HOF-DEMO-XREVIEW",
            "tenant_id": "tenant-demo",
            "procedure_id": "PRC-DEMO-XREVIEW",
            "lot_id": "LOT-DEMO-XREVIEW",
            "status": "published",
            "immutable": True,
            "procedure": {
                "id": "PRC-DEMO-XREVIEW",
                "title": "Servicio de soporte a usuarios",
                "file_number": "2026/XRV/001",
                "contracting_body": "Ayuntamiento demo",
                "responsible_unit": "Sistemas",
            },
            "lot": {"id": "LOT-DEMO-XREVIEW", "code": "L1", "title": "Soporte de primer nivel", "budget_without_tax": 72000},
            "submission_id": "SUB-DEMO-XREVIEW",
            "bidder": {"id": "BID-DEMO-XREVIEW", "name": "Proveedor adjudicatario demo", "tax_id": "B00000000"},
            "awarded_amount": 68400,
            "result": {"total_score": "91.50", "eligible": True},
            "decision_reason": "Acuerdo validado por la mesa",
            "documents": [{
                "id": "DOC-DEMO-XREVIEW", "original_filename": "oferta-adjudicataria.pdf",
                "content_hash": "demo-xreview-content-hash", "object_key": "xreview/demo/oferta-adjudicataria.pdf",
                "content_type": "application/pdf", "extracted_text": "Compromiso de atención en menos de cuatro horas.",
                "envelope_kind": "tecnico",
            }],
            "base_documents": [{
                "id": "DOC-BASE-DEMO-XREVIEW", "original_filename": "pcap.pdf",
                "content_hash": "demo-xreview-base-hash", "object_key": "xreview/demo/pcap.pdf",
                "content_type": "application/pdf", "extracted_text": "Condiciones de ejecución del contrato.",
                "envelope_kind": "bases", "document_role": "pcap",
            }],
            "accepted_commitments": [{
                "id": "ASM-DEMO-XREVIEW", "rule_id": "RULE-DEMO-XREVIEW", "result": "cumple",
                "reasoning": "Atención de incidencias críticas en un máximo de cuatro horas.", "citations": [],
            }],
        }
    },
}

_DATA_PROTECTION_RULES: list[dict[str, Any]] = [
    {"kind": "personal_data", "label": "email", "severity": "media", "pattern": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"},
    {"kind": "personal_data", "label": "dni_nie", "severity": "alta", "pattern": r"\b(?:\d{8}[A-HJ-NP-TV-Z]|[XYZ]\d{7}[A-Z])\b"},
    {"kind": "personal_data", "label": "telefono", "severity": "media", "pattern": r"(?<!\d)(?:\+34[\s-]?)?[6789]\d{2}[\s-]?\d{3}[\s-]?\d{3}(?!\d)"},
    {"kind": "personal_data", "label": "iban", "severity": "alta", "pattern": r"\bES\d{2}[\s-]?\d{4}[\s-]?\d{4}[\s-]?\d{2}[\s-]?\d{10}\b"},
    {"kind": "confidential", "label": "confidencialidad", "severity": "alta", "pattern": r"\b(confidencial|reservado|secreto empresarial|secretos empresariales)\b"},
    {"kind": "confidential", "label": "secreto_tecnico", "severity": "critica", "pattern": r"\b(api[_ -]?key|token|password|contrase(?:ñ|n)a|clave privada)\b"},
]

_SEVERITY_RANK = {"baja": 0, "media": 1, "alta": 2, "critica": 3}
_EVIDENCE_TYPES = {"document", "winning_offer", "tender_basis", "contract", "specification", "award", "image", "file"}


def trace_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _add_days(date_value: str | None, days: int) -> str | None:
    if not date_value:
        return None
    try:
        return (date.fromisoformat(date_value[:10]) + timedelta(days=days)).isoformat()
    except ValueError:
        return None


def _add_months(date_value: str | None, months: int | None) -> str | None:
    if not date_value or not months:
        return None
    try:
        parsed = date.fromisoformat(date_value[:10])
    except ValueError:
        return None
    month_index = parsed.month - 1 + months
    year = parsed.year + month_index // 12
    month = month_index % 12 + 1
    day = min(parsed.day, monthrange(year, month)[1])
    return date(year, month, day).isoformat()


def _amount(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Decimal):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def reset_memory() -> None:
    for key in ["contracts", "obligations", "alerts", "evidence", "checks", "documents", "penalties", "invoices", "proceedings", "changes", "feedback", "data_protection", "tasks", "comments", "alert_rules", "audit", "evidence_files"]:
        _MEMORY[key].clear()
    db.reset_db_probe()


def _id(prefix: str) -> str:
    return f"{prefix}-{datetime.now(timezone.utc).strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}"


def _json_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    return value


def _row(row: dict[str, Any] | None) -> dict[str, Any] | None:
    return _json_safe(row) if row else None


def _audit(user: UserContext, action: str, contract_id: str | None, payload: dict[str, Any] | None = None) -> str:
    event_id = trace_id("audit")
    payload = payload or {}
    event = {
        "id": event_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "user_id": user.user_id,
        "action": action,
        "trace_id": event_id,
        "payload": payload,
        "created_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.audit_events(id, tenant_id, contract_id, user_id, action, trace_id, payload)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                (event_id, user.tenant_id, contract_id, user.user_id, action, event_id, Json(payload)),
            )
            conn.commit()
    else:
        _MEMORY["audit"][event_id] = event
    return event_id


def _manager(user: UserContext) -> bool:
    return bool({"administrador", "admin", "responsable_contratacion", "responsable_contrato"} & set(user.roles))


def _contract_visible(user: UserContext, contract: dict[str, Any]) -> bool:
    return contract.get("tenant_id") == user.tenant_id


def _normalize_contract(row: dict[str, Any]) -> dict[str, Any]:
    contract = _json_safe(row)
    contract["cpv_codes"] = contract.get("cpv_codes") or []
    contract["data"] = contract.get("data") or {}
    contract["archived"] = bool(contract.get("archived_at")) or contract.get("status") == "archived"
    contract["counts"] = _contract_counts(contract["tenant_id"], contract["id"])
    return contract


def _contract_counts(tenant_id: str, contract_id: str) -> dict[str, int]:
    if db.db_available():
        row = db.fetch_one(
            """
            SELECT
              (SELECT count(*) FROM xfollow.contract_obligations WHERE tenant_id = %s AND contract_id = %s) AS obligations,
              (SELECT count(*) FROM xfollow.contract_alerts WHERE tenant_id = %s AND contract_id = %s AND status = 'open') AS open_alerts,
              (SELECT count(*) FROM xfollow.contract_evidence WHERE tenant_id = %s AND contract_id = %s) AS evidence,
              (SELECT count(*) FROM xfollow.compliance_checks WHERE tenant_id = %s AND contract_id = %s) AS checks,
              (SELECT count(*) FROM xfollow.penalty_cases WHERE tenant_id = %s AND contract_id = %s) AS penalties,
              (SELECT count(*) FROM xfollow.contract_invoices WHERE tenant_id = %s AND contract_id = %s AND status NOT IN ('rejected')) AS invoices,
              (SELECT count(*) FROM xfollow.contract_invoices WHERE tenant_id = %s AND contract_id = %s AND status IN ('registered', 'conforming', 'payment_ordered')) AS pending_invoices,
              (SELECT count(*) FROM xfollow.contract_proceedings WHERE tenant_id = %s AND contract_id = %s AND status NOT IN ('resolved', 'closed', 'rejected')) AS open_proceedings,
              (SELECT count(*) FROM xfollow.contract_changes WHERE tenant_id = %s AND contract_id = %s AND status NOT IN ('applied', 'rejected')) AS open_changes,
              (SELECT count(*) FROM xfollow.data_protection_assessments WHERE tenant_id = %s AND contract_id = %s AND status = 'draft' AND result <> 'clear') AS data_protection_open,
              (SELECT count(*) FROM xfollow.contract_tasks WHERE tenant_id = %s AND contract_id = %s AND status IN ('todo', 'doing')) AS open_tasks
            """,
            (
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
                tenant_id,
                contract_id,
            ),
        )
        return {key: int(row[key] or 0) for key in row} if row else {}
    return {
        "obligations": sum(1 for item in _MEMORY["obligations"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id),
        "open_alerts": sum(1 for item in _MEMORY["alerts"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id and item["status"] == "open"),
        "evidence": sum(1 for item in _MEMORY["evidence"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id),
        "checks": sum(1 for item in _MEMORY["checks"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id),
        "penalties": sum(1 for item in _MEMORY["penalties"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id),
        "invoices": sum(1 for item in _MEMORY["invoices"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id and item["status"] != "rejected"),
        "pending_invoices": sum(1 for item in _MEMORY["invoices"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id and item["status"] in {"registered", "conforming", "payment_ordered"}),
        "open_proceedings": sum(1 for item in _MEMORY["proceedings"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id and item["status"] not in {"resolved", "closed", "rejected"}),
        "open_changes": sum(1 for item in _MEMORY["changes"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id and item["status"] not in {"applied", "rejected"}),
        "data_protection_open": sum(1 for item in _MEMORY["data_protection"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id and item["status"] == "draft" and item["result"] != "clear"),
        "open_tasks": sum(1 for item in _MEMORY["tasks"].values() if item["tenant_id"] == tenant_id and item["contract_id"] == contract_id and item["status"] in {"todo", "doing"}),
    }


def list_contracts(
    user: UserContext,
    include_archived: bool = False,
    query: str | None = None,
    status: str | None = None,
    renewal_status: str | None = None,
    responsible_unit: str | None = None,
    contractor: str | None = None,
) -> dict[str, Any]:
    if db.db_available():
        clauses = ["tenant_id = %s"]
        params: list[Any] = [user.tenant_id]
        if not include_archived:
            clauses.append("status <> 'archived' AND archived_at IS NULL")
        if query:
            clauses.append("(title ILIKE %s OR coalesce(file_number, '') ILIKE %s OR coalesce(contractor, '') ILIKE %s)")
            params.extend([f"%{query}%", f"%{query}%", f"%{query}%"])
        if status:
            clauses.append("status = %s")
            params.append(status)
        if renewal_status:
            clauses.append("renewal_status = %s")
            params.append(renewal_status)
        if responsible_unit:
            clauses.append("responsible_unit ILIKE %s")
            params.append(f"%{responsible_unit}%")
        if contractor:
            clauses.append("contractor ILIKE %s")
            params.append(f"%{contractor}%")
        rows = db.fetch_all(
            f"SELECT * FROM xfollow.contracts WHERE {' AND '.join(clauses)} ORDER BY updated_at DESC, id",
            tuple(params),
        )
        items = [_normalize_contract(row) for row in rows]
    else:
        items = [
            _normalize_contract(item)
            for item in _MEMORY["contracts"].values()
            if _contract_visible(user, item)
            and (include_archived or (item.get("status") != "archived" and not item.get("archived_at")))
            and (
                not query
                or query.lower()
                in " ".join(
                    str(item.get(field) or "")
                    for field in ("title", "file_number", "contractor")
                ).lower()
            )
            and (not status or item.get("status") == status)
            and (not renewal_status or item.get("renewal_status") == renewal_status)
            and (not responsible_unit or responsible_unit.lower() in (item.get("responsible_unit") or "").lower())
            and (not contractor or contractor.lower() in (item.get("contractor") or "").lower())
        ]
        items.sort(key=lambda item: item.get("updated_at") or "", reverse=True)
    return {"trace_id": trace_id("ctr"), "items": items, "summary": _summary(user)}


def list_contract_comments(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_comments WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["comments"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
        items.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("cmt"), "items": items}


def create_contract_comment(user: UserContext, contract_id: str, payload: ContractCommentCreate) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    comment_id = _id("CMT")
    item = {
        "id": comment_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "body": payload.body,
        "created_by": user.user_id,
        "created_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                "INSERT INTO xfollow.contract_comments(id, tenant_id, contract_id, body, created_by) VALUES (%s, %s, %s, %s, %s)",
                (comment_id, user.tenant_id, contract_id, payload.body, user.user_id),
            )
            conn.commit()
        result = db.fetch_one("SELECT * FROM xfollow.contract_comments WHERE tenant_id = %s AND id = %s", (user.tenant_id, comment_id))
        item = _row(result) or item
    else:
        _MEMORY["comments"][comment_id] = item
    _audit(user, "contract.comment.created", contract_id, {"comment_id": comment_id})
    return item


def list_alert_rules(user: UserContext) -> dict[str, Any]:
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_alert_rules WHERE tenant_id = %s ORDER BY updated_at DESC",
            (user.tenant_id,),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["alert_rules"].values() if item["tenant_id"] == user.tenant_id]
    return {"trace_id": trace_id("alr"), "items": items}


def create_alert_rule(user: UserContext, payload: AlertRuleCreate) -> dict[str, Any]:
    rule_id = _id("ALR")
    item = {
        "id": rule_id,
        "tenant_id": user.tenant_id,
        "name": payload.name,
        "filters": payload.filters,
        "lead_days": payload.lead_days,
        "enabled": True,
        "shared": payload.shared,
        "created_by": user.user_id,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_alert_rules(id, tenant_id, name, filters, lead_days, shared, created_by)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                (rule_id, user.tenant_id, payload.name, Json(payload.filters), payload.lead_days, payload.shared, user.user_id),
            )
            conn.commit()
        result = db.fetch_one("SELECT * FROM xfollow.contract_alert_rules WHERE tenant_id = %s AND id = %s", (user.tenant_id, rule_id))
        item = _row(result) or item
    else:
        _MEMORY["alert_rules"][rule_id] = item
    _audit(user, "alert_rule.created", None, {"alert_rule_id": rule_id})
    return item


def _summary(user: UserContext) -> dict[str, Any]:
    if db.db_available():
        row = db.fetch_one(
            """
            SELECT
              count(*) AS total,
              count(*) FILTER (WHERE status = 'active') AS active,
              count(*) FILTER (WHERE status = 'archived' OR archived_at IS NOT NULL) AS archived,
              (SELECT count(*) FROM xfollow.contract_alerts WHERE tenant_id = %s AND status = 'open') AS open_alerts,
              (SELECT count(*) FROM xfollow.contract_obligations WHERE tenant_id = %s AND validation_status = 'draft') AS draft_obligations,
              (SELECT count(*) FROM xfollow.contract_invoices WHERE tenant_id = %s AND status IN ('registered', 'conforming', 'payment_ordered')) AS pending_invoices,
              (SELECT count(*) FROM xfollow.contract_proceedings WHERE tenant_id = %s AND status NOT IN ('resolved', 'closed', 'rejected')) AS open_proceedings,
              (SELECT count(*) FROM xfollow.contract_changes WHERE tenant_id = %s AND status NOT IN ('applied', 'rejected')) AS open_changes,
              (SELECT count(*) FROM xfollow.data_protection_assessments WHERE tenant_id = %s AND status = 'draft' AND result <> 'clear') AS data_protection_open,
              (SELECT count(*) FROM xfollow.contract_tasks WHERE tenant_id = %s AND status IN ('todo', 'doing')) AS open_tasks
            FROM xfollow.contracts
            WHERE tenant_id = %s
            """,
            (user.tenant_id, user.tenant_id, user.tenant_id, user.tenant_id, user.tenant_id, user.tenant_id, user.tenant_id, user.tenant_id),
        )
        return {key: int(value or 0) for key, value in (row or {}).items()}
    visible = [item for item in _MEMORY["contracts"].values() if _contract_visible(user, item)]
    return {
        "total": len(visible),
        "active": sum(1 for item in visible if item.get("status") == "active"),
        "archived": sum(1 for item in visible if item.get("status") == "archived" or item.get("archived_at")),
        "open_alerts": sum(1 for item in _MEMORY["alerts"].values() if item["tenant_id"] == user.tenant_id and item["status"] == "open"),
        "draft_obligations": sum(1 for item in _MEMORY["obligations"].values() if item["tenant_id"] == user.tenant_id and item["validation_status"] == "draft"),
        "pending_invoices": sum(1 for item in _MEMORY["invoices"].values() if item["tenant_id"] == user.tenant_id and item["status"] in {"registered", "conforming", "payment_ordered"}),
        "open_proceedings": sum(1 for item in _MEMORY["proceedings"].values() if item["tenant_id"] == user.tenant_id and item["status"] not in {"resolved", "closed", "rejected"}),
        "open_changes": sum(1 for item in _MEMORY["changes"].values() if item["tenant_id"] == user.tenant_id and item["status"] not in {"applied", "rejected"}),
        "data_protection_open": sum(1 for item in _MEMORY["data_protection"].values() if item["tenant_id"] == user.tenant_id and item["status"] == "draft" and item["result"] != "clear"),
        "open_tasks": sum(1 for item in _MEMORY["tasks"].values() if item["tenant_id"] == user.tenant_id and item["status"] in {"todo", "doing"}),
    }


def get_contract(user: UserContext, contract_id: str) -> dict[str, Any] | None:
    if db.db_available():
        row = db.fetch_one("SELECT * FROM xfollow.contracts WHERE tenant_id = %s AND id = %s", (user.tenant_id, contract_id))
        return _normalize_contract(row) if row else None
    contract = _MEMORY["contracts"].get(contract_id)
    return _normalize_contract(contract) if contract and _contract_visible(user, contract) else None


def create_contract(
    user: UserContext,
    payload: ContractCreate,
    *,
    source_kind: str = "manual",
    source_id: str | None = None,
    xtender_workspace_id: str | None = None,
    pcsp_tender_id: str | None = None,
    xreview_procedure_id: str | None = None,
    xreview_lot_id: str | None = None,
    xreview_submission_id: str | None = None,
) -> dict[str, Any]:
    contract_id = _id("CTR")
    contract = {
        "id": contract_id,
        "tenant_id": user.tenant_id,
        "source_kind": source_kind,
        "source_id": source_id,
        "xtender_workspace_id": xtender_workspace_id,
        "pcsp_tender_id": pcsp_tender_id,
        "xreview_procedure_id": xreview_procedure_id,
        "xreview_lot_id": xreview_lot_id,
        "xreview_submission_id": xreview_submission_id,
        "created_by": user.user_id,
        "updated_by": user.user_id,
        "created_at": now_iso(),
        "updated_at": now_iso(),
        "archived_at": None,
        **payload.model_dump(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contracts(
                  id, tenant_id, source_kind, source_id, xtender_workspace_id, pcsp_tender_id,
                  xreview_procedure_id, xreview_lot_id, xreview_submission_id,
                  file_number, title, contracting_body, contractor, responsible_unit, contract_manager,
                  contract_type, procedure, cpv_codes, status, budget_without_tax, budget_with_tax,
                  awarded_amount, currency, start_date, end_date, warranty_end_date, extension_deadline,
                  guarantee_amount, renewal_status, decision_due_date, tags, data, created_by, updated_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    contract_id,
                    user.tenant_id,
                    source_kind,
                    source_id,
                    xtender_workspace_id,
                    pcsp_tender_id,
                    xreview_procedure_id,
                    xreview_lot_id,
                    xreview_submission_id,
                    payload.file_number,
                    payload.title,
                    payload.contracting_body,
                    payload.contractor,
                    payload.responsible_unit,
                    payload.contract_manager,
                    payload.contract_type,
                    payload.procedure,
                    Json(payload.cpv_codes),
                    payload.status,
                    payload.budget_without_tax,
                    payload.budget_with_tax,
                    payload.awarded_amount,
                    payload.currency,
                    payload.start_date,
                    payload.end_date,
                    payload.warranty_end_date,
                    payload.extension_deadline,
                    payload.guarantee_amount,
                    payload.renewal_status,
                    payload.decision_due_date,
                    Json(payload.tags),
                    Json(payload.data),
                    user.user_id,
                    user.user_id,
                ),
            )
            cursor.execute(
                """
                INSERT INTO xfollow.contract_members(contract_id, user_id, access_level)
                VALUES (%s, %s, 'manage')
                ON CONFLICT (contract_id, user_id) DO NOTHING
                """,
                (contract_id, user.user_id),
            )
            conn.commit()
    else:
        _MEMORY["contracts"][contract_id] = contract
    _audit(user, "contract.created", contract_id, {"source_kind": source_kind, "source_id": source_id})
    return get_contract(user, contract_id) or contract


def update_contract(user: UserContext, contract_id: str, payload: ContractUpdate) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    updates = payload.model_dump(exclude_unset=True)
    if db.db_available():
        merged = {**contract, **updates}
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contracts
                SET file_number = %s, title = %s, contracting_body = %s, contractor = %s,
                    responsible_unit = %s, contract_manager = %s, contract_type = %s, procedure = %s,
                    cpv_codes = %s, status = %s, budget_without_tax = %s, budget_with_tax = %s,
                    awarded_amount = %s, currency = %s, start_date = %s, end_date = %s,
                    warranty_end_date = %s, extension_deadline = %s, guarantee_amount = %s,
                    renewal_status = %s, decision_due_date = %s, tags = %s,
                    data = %s, updated_by = %s, updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (
                    merged.get("file_number"),
                    merged["title"],
                    merged.get("contracting_body"),
                    merged.get("contractor"),
                    merged.get("responsible_unit"),
                    merged.get("contract_manager"),
                    merged.get("contract_type"),
                    merged.get("procedure"),
                    Json(merged.get("cpv_codes") or []),
                    merged.get("status"),
                    merged.get("budget_without_tax"),
                    merged.get("budget_with_tax"),
                    merged.get("awarded_amount"),
                    merged.get("currency") or "EUR",
                    merged.get("start_date"),
                    merged.get("end_date"),
                    merged.get("warranty_end_date"),
                    merged.get("extension_deadline"),
                    merged.get("guarantee_amount"),
                    merged.get("renewal_status") or "undefined",
                    merged.get("decision_due_date"),
                    Json(merged.get("tags") or []),
                    Json(merged.get("data") or {}),
                    user.user_id,
                    user.tenant_id,
                    contract_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["contracts"][contract_id].update(updates)
        _MEMORY["contracts"][contract_id]["updated_by"] = user.user_id
        _MEMORY["contracts"][contract_id]["updated_at"] = now_iso()
    _audit(user, "contract.updated", contract_id, {"fields": sorted(updates)})
    return get_contract(user, contract_id) or contract


def archive_contract(user: UserContext, contract_id: str, archived: bool) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    status = "archived" if archived else "active"
    archived_at = now_iso() if archived else None
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                "UPDATE xfollow.contracts SET status = %s, archived_at = %s, updated_by = %s, updated_at = now() WHERE tenant_id = %s AND id = %s",
                (status, archived_at, user.user_id, user.tenant_id, contract_id),
            )
            conn.commit()
    else:
        _MEMORY["contracts"][contract_id]["status"] = status
        _MEMORY["contracts"][contract_id]["archived_at"] = archived_at
        _MEMORY["contracts"][contract_id]["updated_at"] = now_iso()
    _audit(user, "contract.archived" if archived else "contract.restored", contract_id)
    return get_contract(user, contract_id) or contract


def xtender_candidates(user: UserContext, query: str | None = None) -> dict[str, Any]:
    if db.db_available():
        items = [_row(item) for item in integrations.xtender_candidates(user.tenant_id, query)]
    else:
        items = [item for item in _MEMORY["xtender"].values() if item["tenant_id"] == user.tenant_id]
    return {"trace_id": trace_id("xte"), "items": items}


def import_xtender(user: UserContext, request: ImportXtenderRequest) -> dict[str, Any]:
    existing = _existing_import(user, "xtender", request.workspace_id)
    if existing:
        return {"trace_id": trace_id("imp"), "contract": existing, "created": False}
    if db.db_available():
        workspace = integrations.get_xtender(user.tenant_id, request.workspace_id)
    else:
        workspace = _MEMORY["xtender"].get(request.workspace_id)
    if not workspace:
        raise KeyError(request.workspace_id)
    data = workspace.get("data") or {}
    payload = ContractCreate(
        title=workspace.get("title") or data.get("title") or "Contrato importado de xTender",
        file_number=data.get("file_number"),
        contracting_body=data.get("contracting_body"),
        responsible_unit=workspace.get("unit") or data.get("promoting_unit"),
        contract_manager=data.get("contract_manager"),
        contract_type=data.get("contract_type"),
        procedure=data.get("procedure"),
        cpv_codes=data.get("cpv_codes") or ([data["cpv"]] if data.get("cpv") else []),
        status="active",
        budget_without_tax=data.get("budget") or data.get("budget_without_tax"),
        budget_with_tax=data.get("budget_with_tax"),
        awarded_amount=data.get("awarded_amount"),
        start_date=data.get("start_date"),
        end_date=data.get("end_date"),
        warranty_end_date=data.get("warranty_end_date"),
        extension_deadline=data.get("extension_deadline"),
        guarantee_amount=data.get("guarantee_amount"),
        data={**data, "imported_from": "xtender", "xtender_workspace_id": request.workspace_id},
    )
    try:
        contract = create_contract(user, payload, source_kind="xtender", source_id=request.workspace_id, xtender_workspace_id=request.workspace_id)
    except UniqueViolation:
        existing = _existing_import(user, "xtender", request.workspace_id)
        if existing:
            return {"trace_id": trace_id("imp"), "contract": existing, "created": False}
        raise
    extracted = extract_obligations(user, contract["id"])["items"]
    return {"trace_id": trace_id("imp"), "contract": contract, "created": True, "draft_obligations": extracted}


def pcsp_candidates(user: UserContext, query: str | None = None) -> dict[str, Any]:
    if db.db_available():
        items = [_row(item) for item in integrations.pcsp_candidates(query)]
    else:
        items = list(_MEMORY["pcsp"].values())
    return {"trace_id": trace_id("pcsp"), "items": items}


def import_pcsp(user: UserContext, request: ImportPcspRequest) -> dict[str, Any]:
    existing = _existing_import(user, "pcsp", request.tender_id)
    if existing:
        return {"trace_id": trace_id("imp"), "contract": existing, "created": False}
    if db.db_available():
        tender = integrations.get_pcsp(request.tender_id)
    else:
        tender = _MEMORY["pcsp"].get(request.tender_id)
    if not tender:
        raise KeyError(request.tender_id)
    payload = ContractCreate(
        title=tender.get("title") or "Contrato PCSP",
        file_number=tender.get("expediente"),
        contracting_body=tender.get("contracting_body"),
        cpv_codes=[tender["cpv"]] if tender.get("cpv") else [],
        status="active",
        budget_without_tax=tender.get("budget_without_tax"),
        budget_with_tax=tender.get("budget_with_tax"),
        awarded_amount=tender.get("budget_without_tax"),
        currency=tender.get("currency") or "EUR",
        data={**_json_safe(tender), "imported_from": "pcsp"},
    )
    try:
        contract = create_contract(user, payload, source_kind="pcsp", source_id=tender["id"], pcsp_tender_id=tender["id"])
    except UniqueViolation:
        existing = _existing_import(user, "pcsp", tender["id"])
        if existing:
            return {"trace_id": trace_id("imp"), "contract": existing, "created": False}
        raise
    extracted = extract_obligations(user, contract["id"])["items"]
    return {"trace_id": trace_id("imp"), "contract": contract, "created": True, "draft_obligations": extracted}


def _xreview_data(row: dict[str, Any]) -> dict[str, Any]:
    data = row.get("data") if isinstance(row.get("data"), dict) else row
    return {
        **data,
        "id": row.get("id") or data.get("id"),
        "procedure_id": row.get("procedure_id") or data.get("procedure_id"),
        "lot_id": row.get("lot_id") or row.get("parent_id") or data.get("lot_id"),
        "status": row.get("status") or data.get("status"),
    }


def xreview_candidates(user: UserContext, query: str | None = None) -> dict[str, Any]:
    if db.db_available():
        rows = integrations.xreview_candidates(user.tenant_id, query)
        items = [_json_safe(_xreview_data(row)) for row in rows]
    else:
        items = [
            _json_safe(item)
            for item in _MEMORY["xreview"].values()
            if item["tenant_id"] == user.tenant_id
            and (not query or query.lower() in json.dumps(item, ensure_ascii=False).lower())
        ]
    for item in items:
        procedure_id = str(item.get("procedure_id") or "")
        lot_id = str(item.get("lot_id") or "")
        item["imported_contract_id"] = (_existing_import(user, "xreview", f"{procedure_id}:{lot_id}") or {}).get("id")
    return {"trace_id": trace_id("xrv"), "items": items}


def _apply_xreview_handoff(user: UserContext, contract: dict[str, Any], handoff: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    existing_evidence = list_evidence(user, contract["id"])["items"]
    evidence_by_hash = {str(item.get("content_hash")): item for item in existing_evidence if item.get("content_hash")}

    def link_documents(items: list[dict[str, Any]], evidence_type: str, default_title: str) -> list[dict[str, Any]]:
        linked: list[dict[str, Any]] = []
        for document in items:
            digest = str(document.get("content_hash") or "")
            if digest and digest in evidence_by_hash:
                linked.append(evidence_by_hash[digest])
                continue
            evidence = create_evidence(
                user,
                contract["id"],
                EvidenceCreate(
                    title=document.get("original_filename") or default_title,
                    evidence_type=evidence_type,
                    original_filename=document.get("original_filename"),
                    object_key=(f"s3://{document.get('object_bucket') or 'xreview'}/{str(document['object_key']).lstrip('/')}" if document.get("object_key") else None),
                    content_hash=document.get("content_hash"),
                    extracted_text=document.get("extracted_text"),
                ),
            )
            linked.append(evidence)
            if evidence.get("content_hash"):
                evidence_by_hash[str(evidence["content_hash"])] = evidence
        return linked

    linked_evidence = link_documents(handoff.get("documents") or [], "winning_offer", "Oferta adjudicataria xreview")
    base_evidence = link_documents(handoff.get("base_documents") or [], "tender_basis", "Bases de la licitación xreview")

    existing_obligations = list_obligations(user, contract["id"])["items"]
    imported_assessment_ids = {
        str(ref.get("source_id"))
        for obligation in existing_obligations
        for ref in obligation.get("source_refs") or []
        if ref.get("source_kind") == "xreview_assessment"
    }
    obligations: list[dict[str, Any]] = []
    for index, commitment in enumerate(handoff.get("accepted_commitments") or [], start=1):
        assessment_id = str(commitment.get("id") or "")
        if assessment_id and assessment_id in imported_assessment_ids:
            existing = next(
                item for item in existing_obligations
                if any(ref.get("source_kind") == "xreview_assessment" and str(ref.get("source_id")) == assessment_id for ref in item.get("source_refs") or [])
            )
            obligations.append(existing)
            continue
        reasoning = str(commitment.get("reasoning") or "Compromiso aceptado en la oferta adjudicataria.").strip()
        rule = commitment.get("rule") if isinstance(commitment.get("rule"), dict) else {}
        obligation = create_obligation(
            user,
            contract["id"],
            ObligationCreate(
                category="special_condition",
                title=str(rule.get("title") or f"Compromiso de la oferta {index}"),
                description=reasoning,
                source_refs=[{
                    "source_kind": "xreview_assessment",
                    "source_id": assessment_id or commitment.get("rule_id"),
                    "rule_id": commitment.get("rule_id"),
                    "rule_code": rule.get("code"),
                    "rule_kind": rule.get("kind"),
                    "handoff_id": handoff.get("id"),
                    "citations": commitment.get("citations") or [],
                }],
                severity="media",
                status="pending",
                validation_status="validated",
            ),
        )
        obligations.append(obligation)
        if assessment_id:
            imported_assessment_ids.add(assessment_id)
    return linked_evidence, base_evidence, obligations


def import_xreview(user: UserContext, request: ImportXreviewRequest) -> dict[str, Any]:
    if db.db_available():
        row = integrations.get_xreview_handoff(user.tenant_id, request.handoff_id)
        handoff = _xreview_data(row) if row else None
    else:
        handoff = _MEMORY["xreview"].get(request.handoff_id)
    if not handoff or handoff.get("status") != "published" or not handoff.get("immutable"):
        raise KeyError(request.handoff_id)

    procedure = handoff.get("procedure") or {}
    lot = handoff.get("lot") or {}
    bidder = handoff.get("bidder") or {}
    procedure_id = str(handoff.get("procedure_id") or procedure.get("id") or "")
    lot_id = str(handoff.get("lot_id") or lot.get("id") or "")
    submission_id = str(handoff.get("submission_id") or "")
    if not procedure_id or not lot_id or not submission_id:
        raise ValueError("invalid_xreview_handoff")
    source_id = f"{procedure_id}:{lot_id}"
    contract = _existing_import(user, "xreview", source_id)
    created = contract is None
    if not contract:
        payload = ContractCreate(
            title=" · ".join(part for part in [procedure.get("title"), lot.get("code"), lot.get("title")] if part) or "Contrato importado de xreview",
            file_number=procedure.get("file_number"),
            contracting_body=procedure.get("contracting_body"),
            contractor=bidder.get("name"),
            responsible_unit=procedure.get("responsible_unit"),
            contract_type=procedure.get("contract_type"),
            procedure=procedure.get("procedure_type"),
            cpv_codes=lot.get("cpv_codes") or procedure.get("cpv_codes") or [],
            status="active",
            budget_without_tax=lot.get("budget_without_tax"),
            awarded_amount=handoff.get("awarded_amount"),
            currency=procedure.get("currency") or "EUR",
            tags=["xreview", f"lote:{lot.get('code') or lot_id}"],
            data={
                "imported_from": "xreview",
                "handoff_id": handoff.get("id"),
                "xreview_procedure_id": procedure_id,
                "xreview_lot_id": lot_id,
                "xreview_submission_id": submission_id,
                "award_result": handoff.get("result"),
                "award_decision_reason": handoff.get("decision_reason"),
                "bidder_tax_id": bidder.get("tax_id"),
            },
        )
        try:
            contract = create_contract(
                user,
                payload,
                source_kind="xreview",
                source_id=source_id,
                xreview_procedure_id=procedure_id,
                xreview_lot_id=lot_id,
                xreview_submission_id=submission_id,
            )
        except UniqueViolation:
            contract = _existing_import(user, "xreview", source_id)
            if not contract:
                raise
            created = False
    evidence, base_evidence, obligations = _apply_xreview_handoff(user, contract, handoff)
    _audit(user, "contract.imported.xreview", contract["id"], {"handoff_id": handoff.get("id"), "created": created, "evidence": len(evidence), "base_evidence": len(base_evidence), "obligations": len(obligations)})
    return {
        "trace_id": trace_id("imp"), "contract": get_contract(user, contract["id"]) or contract,
        "created": created, "winning_offer_evidence": evidence, "basis_evidence": base_evidence, "accepted_obligations": obligations,
    }


def _existing_import(user: UserContext, source_kind: str, source_id: str) -> dict[str, Any] | None:
    if db.db_available():
        row = db.fetch_one(
            "SELECT * FROM xfollow.contracts WHERE tenant_id = %s AND source_kind = %s AND source_id = %s",
            (user.tenant_id, source_kind, source_id),
        )
        return _normalize_contract(row) if row else None
    for contract in _MEMORY["contracts"].values():
        if contract["tenant_id"] == user.tenant_id and contract.get("source_kind") == source_kind and contract.get("source_id") == source_id:
            return _normalize_contract(contract)
    return None


def list_obligations(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_obligations WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC, id",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["obligations"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("obl"), "items": items}


def create_obligation(user: UserContext, contract_id: str, payload: ObligationCreate) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    obligation_id = _id("OBL")
    item = {
        "id": obligation_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "created_by": user.user_id,
        "validated_by": user.user_id if payload.validation_status == "validated" else None,
        "validated_at": now_iso() if payload.validation_status == "validated" else None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
        **payload.model_dump(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_obligations(
                  id, tenant_id, contract_id, category, title, description, source_refs, due_date,
                  recurrence, severity, status, validation_status, created_by, validated_by, validated_at
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    obligation_id,
                    user.tenant_id,
                    contract_id,
                    payload.category,
                    payload.title,
                    payload.description,
                    Json(payload.source_refs),
                    payload.due_date,
                    payload.recurrence,
                    payload.severity,
                    payload.status,
                    payload.validation_status,
                    user.user_id,
                    user.user_id if payload.validation_status == "validated" else None,
                    now_iso() if payload.validation_status == "validated" else None,
                ),
            )
            conn.commit()
    else:
        _MEMORY["obligations"][obligation_id] = item
    if payload.due_date:
        _ensure_alert_for_obligation(user, contract_id, item)
    _audit(user, "obligation.created", contract_id, {"obligation_id": obligation_id})
    return get_obligation(user, obligation_id) or item


def get_obligation(user: UserContext, obligation_id: str) -> dict[str, Any] | None:
    if db.db_available():
        row = db.fetch_one("SELECT * FROM xfollow.contract_obligations WHERE tenant_id = %s AND id = %s", (user.tenant_id, obligation_id))
        return _row(row)
    item = _MEMORY["obligations"].get(obligation_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def update_obligation(user: UserContext, obligation_id: str, payload: ObligationUpdate) -> dict[str, Any]:
    item = get_obligation(user, obligation_id)
    if not item:
        raise KeyError(obligation_id)
    updates = payload.model_dump(exclude_unset=True)
    merged = {**item, **updates}
    if updates.get("validation_status") == "validated":
        merged["validated_by"] = user.user_id
        merged["validated_at"] = now_iso()
        if merged.get("status") == "draft":
            merged["status"] = "pending"
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contract_obligations
                SET category = %s, title = %s, description = %s, source_refs = %s, due_date = %s,
                    recurrence = %s, severity = %s, status = %s, validation_status = %s,
                    validated_by = %s, validated_at = %s, updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (
                    merged["category"],
                    merged["title"],
                    merged["description"],
                    Json(merged.get("source_refs") or []),
                    merged.get("due_date"),
                    merged.get("recurrence"),
                    merged.get("severity"),
                    merged.get("status"),
                    merged.get("validation_status"),
                    merged.get("validated_by"),
                    merged.get("validated_at"),
                    user.tenant_id,
                    obligation_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["obligations"][obligation_id].update(merged)
        _MEMORY["obligations"][obligation_id]["updated_at"] = now_iso()
    if merged.get("due_date") and merged.get("validation_status") == "validated":
        _ensure_alert_for_obligation(user, merged["contract_id"], merged)
    _audit(user, "obligation.updated", merged["contract_id"], {"obligation_id": obligation_id, "fields": sorted(updates)})
    return get_obligation(user, obligation_id) or merged


def extract_obligations(user: UserContext, contract_id: str) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    source_ref = {"source_kind": contract.get("source_kind"), "source_id": contract.get("source_id"), "contract_id": contract_id}
    proposals: list[ObligationCreate] = []
    if contract.get("end_date"):
        proposals.append(ObligationCreate(category="deadline", title="Fin de ejecucion contractual", description="Controlar la fecha final de ejecucion y preparar la recepcion o continuidad del servicio.", due_date=contract["end_date"], severity="alta", source_refs=[source_ref]))
    if contract.get("extension_deadline") or contract.get("data", {}).get("extensions"):
        proposals.append(ObligationCreate(category="extension", title="Preaviso y decision de prorroga", description=str(contract.get("data", {}).get("extensions") or "Revisar la posible prorroga y su preaviso antes de vencer el contrato."), due_date=contract.get("extension_deadline") or contract.get("end_date"), severity="media", source_refs=[source_ref]))
    if contract.get("warranty_end_date"):
        proposals.append(ObligationCreate(category="warranty", title="Fin del plazo de garantia", description="Verificar incidencias abiertas antes de finalizar la garantia.", due_date=contract["warranty_end_date"], severity="media", source_refs=[source_ref]))
    for raw in contract.get("data", {}).get("milestones") or []:
        if isinstance(raw, dict):
            proposals.append(ObligationCreate(category="delivery", title=str(raw.get("title") or raw.get("name") or "Hito de ejecucion"), description=str(raw.get("description") or raw.get("notes") or "Hito importado del expediente."), due_date=raw.get("date") or raw.get("due_date"), severity="media", source_refs=[source_ref]))
    for raw in contract.get("data", {}).get("special_execution_conditions") or []:
        if isinstance(raw, dict):
            proposals.append(ObligationCreate(category="special_condition", title=str(raw.get("title") or "Condicion especial de ejecucion"), description=str(raw.get("description") or raw.get("text") or raw), severity="alta", source_refs=[source_ref]))
        elif raw:
            proposals.append(ObligationCreate(category="special_condition", title="Condicion especial de ejecucion", description=str(raw), severity="alta", source_refs=[source_ref]))
    proposals.extend(_textual_obligation_proposals(contract, _source_fragments_for_contract(contract)))
    if not proposals:
        proposals.append(ObligationCreate(category="documentation", title="Revisar obligaciones de ejecucion", description="No constan hitos estructurados. Revisar PPT, PCAP, contrato formalizado y resolucion de adjudicacion para validar obligaciones de seguimiento.", severity="media", source_refs=[source_ref]))

    existing_titles = {item["title"].lower() for item in list_obligations(user, contract_id)["items"]}
    created = [create_obligation(user, contract_id, proposal) for proposal in proposals if proposal.title.lower() not in existing_titles]
    _audit(user, "obligations.extracted", contract_id, {"created": len(created)})
    return {"trace_id": trace_id("obl"), "items": created}


def _source_fragments_for_contract(contract: dict[str, Any]) -> list[dict[str, Any]]:
    fragments: list[dict[str, Any]] = []
    data = contract.get("data") or {}
    for key in [
        "object",
        "need",
        "included",
        "excluded",
        "outcome",
        "duration",
        "extensions",
        "milestones",
        "service_levels",
        "deliverables",
        "special_execution_conditions",
        "penalties",
        "security",
    ]:
        value = data.get(key)
        if value:
            fragments.append({"source_kind": "xtender_data", "source_id": contract.get("xtender_workspace_id") or contract["id"], "section": key, "text": _stringify_source_value(value)})
    if db.db_available() and contract.get("pcsp_tender_id"):
        rows = integrations.pcsp_execution_fragments(contract["pcsp_tender_id"])
        for row in rows:
            fragments.append(
                {
                    "source_kind": "pcsp_chunk",
                    "source_id": row["id"],
                    "document_id": row["document_id"],
                    "document_type": row["document_type"],
                    "title": row["title"],
                    "source_url": row["source_url"],
                    "heading_path": row.get("heading_path") or [],
                    "page_start": row.get("page_start"),
                    "text": row["chunk_text"],
                }
            )
    return fragments


def _stringify_source_value(value: Any) -> str:
    if isinstance(value, str):
        return value
    return json.dumps(_json_safe(value), ensure_ascii=False)


def _textual_obligation_proposals(contract: dict[str, Any], fragments: list[dict[str, Any]]) -> list[ObligationCreate]:
    proposals: list[ObligationCreate] = []
    seen: set[str] = set()
    rules: list[tuple[str, str, str, str]] = [
        ("extension", "prorroga", "Preaviso y decision de prorroga", "Revisar la posible prorroga, su preaviso y la documentacion necesaria."),
        ("warranty", "garant", "Seguimiento del plazo de garantia", "Controlar el plazo de garantia, incidencias pendientes y devolucion de garantia cuando proceda."),
        ("deadline", "plazo", "Control de plazos de ejecucion", "Verificar los plazos parciales y totales de ejecucion contractual."),
        ("delivery", "entreg", "Entregables y recepcion", "Comprobar entregables, criterios de aceptacion y recepcion formal."),
        ("payment", "factur", "Facturacion y conformidad", "Revisar facturas, conformidad y pagos asociados a la ejecucion."),
        ("sla", "nivel de servicio", "Niveles de servicio", "Medir indicadores, SLA y consecuencias asociadas al incumplimiento."),
        ("special_condition", "condicion especial", "Condiciones especiales de ejecucion", "Acreditar las condiciones especiales de ejecucion previstas en el contrato."),
        ("other", "penal", "Penalidades por incumplimiento", "Preparar el calculo y la tramitacion de penalidades cuando existan incumplimientos."),
        ("documentation", "recepci", "Acta de recepcion", "Preparar acta de recepcion y evidencias de conformidad."),
        ("documentation", "liquidaci", "Informe de liquidacion", "Preparar informe de liquidacion y cierre economico del contrato."),
        ("other", "resoluci", "Riesgo de resolucion contractual", "Revisar posibles causas de resolucion y documentar actuaciones."),
    ]
    for fragment in fragments:
        text = str(fragment.get("text") or "")
        normalized = _normalize_text(text)
        if not normalized:
            continue
        for category, needle, title, description in rules:
            if needle not in normalized:
                continue
            evidence_sentence = _best_sentence(text, needle)
            key = f"{title}:{evidence_sentence[:80].lower()}"
            if key in seen:
                continue
            seen.add(key)
            proposals.append(
                ObligationCreate(
                    category=category,  # type: ignore[arg-type]
                    title=title,
                    description=f"{description}\n\nFuente: {evidence_sentence}"[:1200],
                    due_date=_date_hint(contract, category),
                    severity="alta" if category in {"other", "special_condition", "deadline"} else "media",
                    source_refs=[_source_ref_from_fragment(fragment, evidence_sentence)],
                )
            )
    return proposals[:24]


def _normalize_text(value: str) -> str:
    text = value.lower()
    replacements = {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u", "à": "a", "è": "e", "ò": "o", "ü": "u"}
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _best_sentence(text: str, needle: str) -> str:
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text)
    normalized_needle = _normalize_text(needle)
    for sentence in sentences:
        clean = " ".join(sentence.split())
        if normalized_needle in _normalize_text(clean):
            return clean[:500]
    return " ".join(text.split())[:500]


def _date_hint(contract: dict[str, Any], category: str) -> str | None:
    if category == "extension":
        return contract.get("extension_deadline") or contract.get("end_date")
    if category == "warranty":
        return contract.get("warranty_end_date")
    if category in {"deadline", "delivery", "documentation"}:
        return contract.get("end_date")
    return None


def _source_ref_from_fragment(fragment: dict[str, Any], quote: str) -> dict[str, Any]:
    return {
        "source_kind": fragment.get("source_kind"),
        "source_id": fragment.get("source_id"),
        "document_id": fragment.get("document_id"),
        "document_type": fragment.get("document_type"),
        "title": fragment.get("title"),
        "source_url": fragment.get("source_url"),
        "heading_path": fragment.get("heading_path") or [],
        "page_start": fragment.get("page_start"),
        "quote": quote,
    }


def _ensure_alert_for_obligation(user: UserContext, contract_id: str, obligation: dict[str, Any]) -> None:
    due_date = obligation.get("due_date")
    if not due_date:
        return
    title = f"{obligation['title']}: vencimiento"
    if db.db_available():
        existing = db.fetch_one(
            "SELECT id FROM xfollow.contract_alerts WHERE tenant_id = %s AND contract_id = %s AND obligation_id = %s AND title = %s",
            (user.tenant_id, contract_id, obligation["id"], title),
        )
        if existing:
            return
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_alerts(id, tenant_id, contract_id, obligation_id, alert_type, title, due_date, lead_days, severity)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (_id("ALT"), user.tenant_id, contract_id, obligation["id"], obligation.get("category") or "deadline", title, due_date, 30, obligation.get("severity") or "media"),
            )
            conn.commit()
    else:
        if any(item["obligation_id"] == obligation["id"] and item["title"] == title for item in _MEMORY["alerts"].values()):
            return
        alert_id = _id("ALT")
        _MEMORY["alerts"][alert_id] = {
            "id": alert_id,
            "tenant_id": user.tenant_id,
            "contract_id": contract_id,
            "obligation_id": obligation["id"],
            "alert_type": obligation.get("category") or "deadline",
            "title": title,
            "due_date": due_date,
            "lead_days": 30,
            "severity": obligation.get("severity") or "media",
            "status": "open",
            "generated_by": "system",
            "created_at": now_iso(),
            "updated_at": now_iso(),
        }


def list_alerts(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all("SELECT * FROM xfollow.contract_alerts WHERE tenant_id = %s AND contract_id = %s ORDER BY due_date, id", (user.tenant_id, contract_id))
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["alerts"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("alt"), "items": items}


def update_alert(user: UserContext, alert_id: str, payload: AlertUpdate) -> dict[str, Any]:
    updates = payload.model_dump(exclude_unset=True)
    if db.db_available():
        current = db.fetch_one("SELECT * FROM xfollow.contract_alerts WHERE tenant_id = %s AND id = %s", (user.tenant_id, alert_id))
        if not current:
            raise KeyError(alert_id)
        merged = {**current, **updates}
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                "UPDATE xfollow.contract_alerts SET status = %s, due_date = %s, lead_days = %s, updated_at = now() WHERE tenant_id = %s AND id = %s",
                (merged["status"], merged["due_date"], merged["lead_days"], user.tenant_id, alert_id),
            )
            conn.commit()
        item = db.fetch_one("SELECT * FROM xfollow.contract_alerts WHERE tenant_id = %s AND id = %s", (user.tenant_id, alert_id))
    else:
        item = _MEMORY["alerts"].get(alert_id)
        if not item or item["tenant_id"] != user.tenant_id:
            raise KeyError(alert_id)
        item.update(updates)
        item["updated_at"] = now_iso()
    _audit(user, "alert.updated", item["contract_id"], {"alert_id": alert_id, "fields": sorted(updates)})
    return _row(item) or item


def list_tasks(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_tasks WHERE tenant_id = %s AND contract_id = %s ORDER BY status, due_date NULLS LAST, created_at DESC",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["tasks"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("tsk"), "items": items}


def create_task(user: UserContext, contract_id: str, payload: TaskCreate) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    task_id = _id("TSK")
    item = {
        "id": task_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "created_by": user.user_id,
        "created_at": now_iso(),
        "updated_at": now_iso(),
        **payload.model_dump(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_tasks(
                  id, tenant_id, contract_id, source_type, source_id, title, description,
                  due_date, owner_user_id, severity, status, created_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    task_id,
                    user.tenant_id,
                    contract_id,
                    payload.source_type,
                    payload.source_id,
                    payload.title,
                    payload.description,
                    payload.due_date,
                    payload.owner_user_id,
                    payload.severity,
                    payload.status,
                    user.user_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["tasks"][task_id] = item
    _audit(user, "task.created", contract_id, {"task_id": task_id, "source_type": payload.source_type, "source_id": payload.source_id})
    return get_task(user, task_id) or item


def get_task(user: UserContext, task_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.contract_tasks WHERE tenant_id = %s AND id = %s", (user.tenant_id, task_id)))
    item = _MEMORY["tasks"].get(task_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def update_task(user: UserContext, task_id: str, payload: TaskUpdate) -> dict[str, Any]:
    current = get_task(user, task_id)
    if not current:
        raise KeyError(task_id)
    updates = payload.model_dump(exclude_unset=True)
    merged = {**current, **updates}
    completed_by = user.user_id if updates.get("status") == "done" else current.get("completed_by")
    completed_at = now_iso() if updates.get("status") == "done" else current.get("completed_at")
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contract_tasks
                SET title = %s, description = %s, due_date = %s, owner_user_id = %s,
                    severity = %s, status = %s, completed_by = %s, completed_at = %s, updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (
                    merged["title"],
                    merged.get("description"),
                    merged.get("due_date"),
                    merged.get("owner_user_id"),
                    merged.get("severity"),
                    merged.get("status"),
                    completed_by,
                    completed_at,
                    user.tenant_id,
                    task_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["tasks"][task_id].update(merged)
        _MEMORY["tasks"][task_id]["completed_by"] = completed_by
        _MEMORY["tasks"][task_id]["completed_at"] = completed_at
        _MEMORY["tasks"][task_id]["updated_at"] = now_iso()
    _audit(user, "task.updated", current["contract_id"], {"task_id": task_id, "fields": sorted(updates)})
    return get_task(user, task_id) or merged


def generate_tasks(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    existing_keys = {
        f"{task.get('source_type')}:{task.get('source_id')}"
        for task in list_tasks(user, contract_id)["items"]
        if task.get("source_type") and task.get("source_id")
    }
    created = []
    for alert in list_alerts(user, contract_id)["items"]:
        key = f"alert:{alert['id']}"
        if alert.get("status") != "open" or key in existing_keys:
            continue
        created.append(
            create_task(
                user,
                contract_id,
                TaskCreate(
                    title=alert["title"],
                    description=f"Gestionar alerta de tipo {alert.get('alert_type')} con preaviso de {alert.get('lead_days')} dias.",
                    due_date=alert.get("due_date"),
                    severity=alert.get("severity") or "media",
                    source_type="alert",
                    source_id=alert["id"],
                ),
            )
        )
    for obligation in list_obligations(user, contract_id)["items"]:
        key = f"obligation:{obligation['id']}"
        if obligation.get("validation_status") != "draft" or key in existing_keys:
            continue
        created.append(
            create_task(
                user,
                contract_id,
                TaskCreate(
                    title=f"Validar: {obligation['title']}",
                    description="Revisar, aceptar, corregir o rechazar la obligacion propuesta antes de usarla en verificaciones.",
                    due_date=obligation.get("due_date"),
                    severity=obligation.get("severity") or "media",
                    source_type="obligation",
                    source_id=obligation["id"],
                ),
            )
        )
    for check in list_checks(user, contract_id)["items"]:
        key = f"check:{check['id']}"
        if check.get("status") != "draft" or key in existing_keys:
            continue
        created.append(
            create_task(
                user,
                contract_id,
                TaskCreate(
                    title=f"Revisar verificacion: {check.get('result')}",
                    description=check.get("reasoning"),
                    severity="alta" if check.get("result") in {"breach", "at_risk"} else "media",
                    source_type="check",
                    source_id=check["id"],
                ),
            )
        )
    for invoice in list_invoices(user, contract_id)["items"]:
        key = f"invoice:{invoice['id']}"
        if invoice.get("status") not in {"registered", "conforming", "payment_ordered"} or key in existing_keys:
            continue
        title = "Conformar factura" if invoice.get("status") == "registered" else "Gestionar pago de factura"
        created.append(
            create_task(
                user,
                contract_id,
                TaskCreate(
                    title=f"{title}: {invoice.get('invoice_number')}",
                    description=invoice.get("concept") or "Revisar factura vinculada a la ejecucion del contrato.",
                    due_date=invoice.get("payment_due_date"),
                    severity="alta" if invoice.get("status") == "payment_ordered" else "media",
                    source_type="invoice",
                    source_id=invoice["id"],
                ),
            )
        )
    for proceeding in list_proceedings(user, contract_id)["items"]:
        key = f"proceeding:{proceeding['id']}"
        if proceeding.get("status") in {"resolved", "closed", "rejected"} or key in existing_keys:
            continue
        created.append(
            create_task(
                user,
                contract_id,
                TaskCreate(
                    title=f"Tramitar expediente: {proceeding.get('title')}",
                    description=proceeding.get("summary"),
                    due_date=proceeding.get("due_date"),
                    severity="alta",
                    source_type="proceeding",
                    source_id=proceeding["id"],
                ),
            )
        )
    for change in list_contract_changes(user, contract_id)["items"]:
        key = f"change:{change['id']}"
        if change.get("status") in {"applied", "rejected"} or key in existing_keys:
            continue
        action = "Validar propuesta" if change.get("status") in {"draft", "under_review"} else "Aplicar propuesta"
        created.append(
            create_task(
                user,
                contract_id,
                TaskCreate(
                    title=f"{action}: {change.get('title')}",
                    description=change.get("impact_summary") or change.get("reason"),
                    due_date=change.get("proposed_start_date") or change.get("proposed_end_date"),
                    severity="alta" if change.get("change_type") == "modification" else "media",
                    source_type="change",
                    source_id=change["id"],
                ),
            )
        )
    for assessment in list_data_protection_assessments(user, contract_id)["items"]:
        key = f"data_protection:{assessment['id']}"
        if assessment.get("status") != "draft" or assessment.get("result") == "clear" or key in existing_keys:
            continue
        created.append(
            create_task(
                user,
                contract_id,
                TaskCreate(
                    title=f"Revisar datos: {assessment.get('source_title')}",
                    description=assessment.get("recommendation"),
                    severity=assessment.get("severity") or "media",
                    source_type="data_protection",
                    source_id=assessment["id"],
                ),
            )
        )
    _audit(user, "tasks.generated", contract_id, {"created": len(created)})
    return {"trace_id": trace_id("tsk"), "items": created}


def create_evidence(user: UserContext, contract_id: str, payload: EvidenceCreate) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if payload.evidence_type not in _EVIDENCE_TYPES:
        raise ValueError("invalid_evidence_type")
    evidence_id = _id("EVD")
    digest = payload.content_hash or hashlib.sha256(f"{payload.title}\n{payload.extracted_text or ''}".encode("utf-8")).hexdigest()
    item = {
        "id": evidence_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "status": "available",
        "uploaded_by": user.user_id,
        "created_at": now_iso(),
        **payload.model_dump(),
        "content_hash": digest,
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_evidence(
                  id, tenant_id, contract_id, obligation_id, title, evidence_type, original_filename,
                  object_key, content_hash, extracted_text, source_url, uploaded_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (evidence_id, user.tenant_id, contract_id, payload.obligation_id, payload.title, payload.evidence_type, payload.original_filename, payload.object_key, digest, payload.extracted_text, payload.source_url, user.user_id),
            )
            conn.commit()
    else:
        _MEMORY["evidence"][evidence_id] = item
    _audit(user, "evidence.created", contract_id, {"evidence_id": evidence_id})
    return _get_evidence(user, evidence_id) or item


def create_evidence_from_upload(
    user: UserContext,
    contract_id: str,
    *,
    filename: str,
    content_type: str | None,
    content: bytes,
    title: str | None = None,
    obligation_id: str | None = None,
    evidence_type: str | None = None,
) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if not content:
        raise ValueError("empty_upload")
    if len(content) > load_runtime_config().upload_max_bytes:
        raise ValueError("file_too_large")
    resolved_type = evidence_type or _evidence_type_for_filename(filename, content_type)
    if resolved_type not in _EVIDENCE_TYPES:
        raise ValueError("invalid_evidence_type")
    digest = hashlib.sha256(content).hexdigest()
    safe_name = _safe_filename(filename or "evidencia.bin")
    object_key = _store_evidence_blob(user.tenant_id, contract_id, digest, safe_name, content)
    try:
        extracted_text = _extract_text_from_bytes(safe_name, content_type, content)
    except Exception as exc:
        raise ValueError("file_text_extraction_failed") from exc
    payload = EvidenceCreate(
        obligation_id=obligation_id,
        title=title or Path(safe_name).stem or "Evidencia",
        evidence_type=resolved_type,
        original_filename=safe_name,
        object_key=object_key,
        content_hash=digest,
        extracted_text=extracted_text,
    )
    evidence = create_evidence(user, contract_id, payload)
    _audit(
        user,
        "evidence.uploaded",
        contract_id,
        {
            "evidence_id": evidence["id"],
            "filename": safe_name,
            "content_type": content_type,
            "evidence_type": resolved_type,
            "bytes": len(content),
            "extracted_chars": len(extracted_text or ""),
        },
    )
    return evidence


def _safe_filename(filename: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", filename.strip())[:180]
    return cleaned or "evidencia.bin"


def _store_evidence_blob(tenant_id: str, contract_id: str, digest: str, filename: str, content: bytes) -> str:
    if db.db_available():
        target = repo_root() / "data" / "evidence" / tenant_id / contract_id
        target.mkdir(parents=True, exist_ok=True)
        path = target / f"{digest[:16]}-{filename}"
        path.write_bytes(content)
        return f"local://data/evidence/{tenant_id}/{contract_id}/{path.name}"
    key = f"memory://{tenant_id}/{contract_id}/{digest[:16]}-{filename}"
    _MEMORY["evidence_files"][key] = {"content": content, "filename": filename}
    return key


def _extract_text_from_bytes(filename: str, content_type: str | None, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".csv", ".json"} or (content_type or "").startswith("text/"):
        return content.decode("utf-8", errors="ignore")[:200000]
    if suffix == ".pdf" or content_type == "application/pdf":
        with fitz.open(stream=content, filetype="pdf") as document:
            return "\n\n".join(page.get_text("text") for page in document)[:300000]
    if suffix == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())[:300000]
    return ""


def _evidence_type_for_filename(filename: str, content_type: str | None) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".pdf", ".docx", ".odt", ".txt", ".md", ".markdown"}:
        return "document"
    if (content_type or "").startswith("image/"):
        return "image"
    return "file"


def _get_evidence(user: UserContext, evidence_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.contract_evidence WHERE tenant_id = %s AND id = %s", (user.tenant_id, evidence_id)))
    item = _MEMORY["evidence"].get(evidence_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def update_evidence(user: UserContext, evidence_id: str, obligation_id: str | None) -> dict[str, Any]:
    evidence = _get_evidence(user, evidence_id)
    if not evidence:
        raise KeyError(evidence_id)
    if obligation_id:
        obligation = get_obligation(user, obligation_id)
        if not obligation or obligation["contract_id"] != evidence["contract_id"]:
            raise ValueError("obligation_not_in_contract")
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                "UPDATE xfollow.contract_evidence SET obligation_id = %s WHERE tenant_id = %s AND id = %s",
                (obligation_id, user.tenant_id, evidence_id),
            )
            conn.commit()
    else:
        _MEMORY["evidence"][evidence_id]["obligation_id"] = obligation_id
    _audit(user, "evidence.linked", evidence["contract_id"], {"evidence_id": evidence_id, "obligation_id": obligation_id})
    return _get_evidence(user, evidence_id) or evidence


def evidence_download_bytes(user: UserContext, evidence_id: str) -> tuple[str, bytes]:
    evidence = _get_evidence(user, evidence_id)
    if not evidence:
        raise KeyError(evidence_id)
    filename = evidence.get("original_filename") or f"{evidence_id}.bin"
    object_key = str(evidence.get("object_key") or "")
    if object_key.startswith("local://data/evidence/"):
        relative = object_key.removeprefix("local://data/evidence/")
        root = (repo_root() / "data" / "evidence").resolve()
        path = (root / relative).resolve()
        if root not in path.parents or not path.is_file():
            raise KeyError(evidence_id)
        return filename, path.read_bytes()
    if object_key.startswith("s3://"):
        bucket_and_key = object_key.removeprefix("s3://")
        bucket, separator, key = bucket_and_key.partition("/")
        if not separator or not bucket or not key:
            raise KeyError(evidence_id)
        config = load_runtime_config()
        try:
            client = boto3.client(
                "s3", endpoint_url=config.object_storage_endpoint,
                aws_access_key_id=config.object_storage_access_key or None,
                aws_secret_access_key=config.object_storage_secret_key or None,
                region_name=config.object_storage_region,
            )
            response = client.get_object(Bucket=bucket, Key=key)
            return filename, response["Body"].read()
        except Exception as exc:
            raise KeyError(evidence_id) from exc
    memory = _MEMORY["evidence_files"].get(object_key)
    if memory:
        return memory.get("filename") or filename, memory["content"]
    raise KeyError(evidence_id)


def list_evidence(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all("SELECT * FROM xfollow.contract_evidence WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC", (user.tenant_id, contract_id))
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["evidence"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("evd"), "items": items}


def list_contract_risks(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_risks WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["risks"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("rsk"), "items": items}


def _get_contract_risk(user: UserContext, risk_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.contract_risks WHERE tenant_id = %s AND id = %s", (user.tenant_id, risk_id)))
    item = _MEMORY["risks"].get(risk_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def create_ai_offer_risks(user: UserContext, contract_id: str, offer_evidence_id: str, result: Any, sources: list[dict[str, Any]]) -> dict[str, Any]:
    offer = _get_evidence(user, offer_evidence_id)
    if not offer or offer["contract_id"] != contract_id or offer.get("evidence_type") != "winning_offer":
        raise ValueError("winning_offer_not_found")
    raw_items = result.get("items") if isinstance(result, dict) else None
    if not isinstance(raw_items, list):
        raise ValueError("ai_offer_risks_invalid_result")
    source_by_id = {str(item["id"]): item for item in sources}
    if offer_evidence_id not in source_by_id:
        raise ValueError("winning_offer_without_extracted_text")
    allowed_types = {"scope", "schedule", "resources", "quality", "economic", "legal", "other"}
    allowed_severity = {"baja", "media", "alta", "critica"}
    created: list[dict[str, Any]] = []
    for raw in raw_items[:24]:
        if not isinstance(raw, dict):
            continue
        title = str(raw.get("title") or "").strip()[:240]
        description = str(raw.get("description") or "").strip()[:5000]
        if not title or not description:
            continue
        risk_type = str(raw.get("risk_type") or "other")
        severity = str(raw.get("severity") or "media")
        if risk_type not in allowed_types:
            risk_type = "other"
        if severity not in allowed_severity:
            severity = "media"
        source_ids = [str(value) for value in raw.get("source_ids") or [] if str(value) in source_by_id]
        if offer_evidence_id not in source_ids:
            source_ids.insert(0, offer_evidence_id)
        citations = [{"evidence_id": source_id, "title": source_by_id[source_id].get("title")} for source_id in source_ids]
        risk_id = _id("RSK")
        item = {
            "id": risk_id,
            "tenant_id": user.tenant_id,
            "contract_id": contract_id,
            "offer_evidence_id": offer_evidence_id,
            "risk_type": risk_type,
            "title": title,
            "description": description,
            "comparison": str(raw.get("comparison") or "Información insuficiente para una comparación concluyente.").strip()[:5000],
            "severity": severity,
            "citations": citations,
            "validation_status": "draft",
            "created_by": user.user_id,
            "created_at": now_iso(),
        }
        if db.db_available():
            with db.connection() as conn, conn.cursor() as cursor:
                cursor.execute(
                    """
                    INSERT INTO xfollow.contract_risks(
                      id, tenant_id, contract_id, offer_evidence_id, risk_type, title, description,
                      comparison, severity, citations, created_by
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (risk_id, user.tenant_id, contract_id, offer_evidence_id, risk_type, title, description, item["comparison"], severity, Json(citations), user.user_id),
                )
                conn.commit()
            item = _get_contract_risk(user, risk_id) or item
        else:
            _MEMORY["risks"][risk_id] = item
        created.append(item)
    _audit(user, "offer.risks_extracted", contract_id, {"offer_evidence_id": offer_evidence_id, "created": len(created)})
    return {"items": created}


def update_contract_risk(user: UserContext, risk_id: str, payload: ContractRiskUpdate) -> dict[str, Any]:
    risk = _get_contract_risk(user, risk_id)
    if not risk:
        raise KeyError(risk_id)
    if not payload.validation_status:
        return risk
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contract_risks
                SET validation_status = %s,
                    validated_by = CASE WHEN %s = 'validated' THEN %s ELSE validated_by END,
                    validated_at = CASE WHEN %s = 'validated' THEN now() ELSE validated_at END,
                    updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (payload.validation_status, payload.validation_status, user.user_id, payload.validation_status, user.tenant_id, risk_id),
            )
            conn.commit()
    else:
        _MEMORY["risks"][risk_id]["validation_status"] = payload.validation_status
        _MEMORY["risks"][risk_id]["validated_by"] = user.user_id if payload.validation_status == "validated" else None
        _MEMORY["risks"][risk_id]["validated_at"] = now_iso() if payload.validation_status == "validated" else None
    _audit(user, "offer.risk_validated", risk["contract_id"], {"risk_id": risk_id, "status": payload.validation_status})
    return _get_contract_risk(user, risk_id) or risk


def list_data_protection_assessments(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.data_protection_assessments WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["data_protection"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
        items.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("dpt"), "items": items}


def scan_data_protection(user: UserContext, contract_id: str) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    sources = [
        {
            "scope": "contract",
            "evidence_id": None,
            "source_title": "Ficha del contrato",
            "text": "\n".join(
                [
                    str(contract.get("title") or ""),
                    str(contract.get("contracting_body") or ""),
                    str(contract.get("contractor") or ""),
                    str(contract.get("contract_manager") or ""),
                    _stringify_source_value(contract.get("data") or {}),
                ]
            ),
        }
    ]
    for evidence in list_evidence(user, contract_id)["items"]:
        sources.append(
            {
                "scope": "evidence",
                "evidence_id": evidence["id"],
                "source_title": evidence.get("title") or evidence.get("original_filename") or evidence["id"],
                "text": "\n".join([str(evidence.get("title") or ""), str(evidence.get("extracted_text") or "")]),
            }
        )
    created = [_insert_data_protection_assessment(user, contract_id, source) for source in sources]
    _audit(user, "data_protection.scanned", contract_id, {"assessments": len(created), "findings": sum(len(item.get("findings") or []) for item in created)})
    return {"trace_id": trace_id("dpt"), "items": created}


def _insert_data_protection_assessment(user: UserContext, contract_id: str, source: dict[str, Any]) -> dict[str, Any]:
    text = str(source.get("text") or "")
    findings = _data_protection_findings(text)
    result = _data_protection_result(findings)
    severity = _max_severity([item["severity"] for item in findings]) if findings else "baja"
    status = "draft" if findings else "resolved"
    assessment_id = _id("DPT")
    item = {
        "id": assessment_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "evidence_id": source.get("evidence_id"),
        "scope": source["scope"],
        "source_title": source["source_title"],
        "result": result,
        "severity": severity,
        "findings": findings,
        "redacted_preview": _redacted_preview(text),
        "recommendation": _data_protection_recommendation(result),
        "status": status,
        "created_by": user.user_id,
        "resolved_by": user.user_id if status == "resolved" else None,
        "resolved_at": now_iso() if status == "resolved" else None,
        "resolution_comment": "Sin hallazgos detectados por reglas locales." if status == "resolved" else None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.data_protection_assessments(
                  id, tenant_id, contract_id, evidence_id, scope, source_title, result,
                  severity, findings, redacted_preview, recommendation, status, created_by,
                  resolved_by, resolved_at, resolution_comment
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    assessment_id,
                    user.tenant_id,
                    contract_id,
                    source.get("evidence_id"),
                    source["scope"],
                    source["source_title"],
                    result,
                    severity,
                    Json(findings),
                    item["redacted_preview"],
                    item["recommendation"],
                    status,
                    user.user_id,
                    item["resolved_by"],
                    item["resolved_at"],
                    item["resolution_comment"],
                ),
            )
            conn.commit()
    else:
        _MEMORY["data_protection"][assessment_id] = item
    return get_data_protection_assessment(user, assessment_id) or item


def get_data_protection_assessment(user: UserContext, assessment_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.data_protection_assessments WHERE tenant_id = %s AND id = %s", (user.tenant_id, assessment_id)))
    item = _MEMORY["data_protection"].get(assessment_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def resolve_data_protection_assessment(user: UserContext, assessment_id: str, payload: DataProtectionAssessmentUpdate) -> dict[str, Any]:
    current = get_data_protection_assessment(user, assessment_id)
    if not current:
        raise KeyError(assessment_id)
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.data_protection_assessments
                SET status = %s, resolved_by = %s, resolved_at = now(),
                    resolution_comment = %s, updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (payload.status, user.user_id, payload.resolution_comment, user.tenant_id, assessment_id),
            )
            conn.commit()
    else:
        _MEMORY["data_protection"][assessment_id]["status"] = payload.status
        _MEMORY["data_protection"][assessment_id]["resolved_by"] = user.user_id
        _MEMORY["data_protection"][assessment_id]["resolved_at"] = now_iso()
        _MEMORY["data_protection"][assessment_id]["resolution_comment"] = payload.resolution_comment
        _MEMORY["data_protection"][assessment_id]["updated_at"] = now_iso()
    _audit(user, "data_protection.resolved", current["contract_id"], {"assessment_id": assessment_id, "status": payload.status})
    return get_data_protection_assessment(user, assessment_id) or current


def _data_protection_findings(text: str) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    for rule in _DATA_PROTECTION_RULES:
        matches = list(re.finditer(rule["pattern"], text, flags=re.IGNORECASE))[:8]
        for match in matches:
            snippet = _snippet(text, match.start(), match.end())
            findings.append(
                {
                    "kind": rule["kind"],
                    "label": rule["label"],
                    "severity": rule["severity"],
                    "match_hash": hashlib.sha256(match.group(0).encode("utf-8", errors="ignore")).hexdigest()[:16],
                    "snippet": _redact_sensitive_text(snippet),
                }
            )
    return findings[:40]


def _data_protection_result(findings: list[dict[str, Any]]) -> str:
    kinds = {item["kind"] for item in findings}
    if not kinds:
        return "clear"
    if kinds == {"personal_data"}:
        return "personal_data"
    if kinds == {"confidential"}:
        return "confidential"
    if kinds:
        return "mixed"
    return "needs_review"


def _data_protection_recommendation(result: str) -> str:
    if result == "clear":
        return "No se detectan datos personales o confidenciales con las reglas locales. Revisar si el documento contiene metadatos o adjuntos no textuales."
    if result == "personal_data":
        return "Aplicar minimizacion: conservar solo datos necesarios, anonimizar o seudonimizar antes de compartir, y revisar base de tratamiento."
    if result == "confidential":
        return "Revisar secreto profesional o informacion estrategica antes de difundir o usar en prompts externos."
    if result == "mixed":
        return "Revisar datos personales y confidencialidad conjuntamente; limitar acceso y documentar decision de tratamiento."
    return "Requiere revision manual antes de reutilizar la evidencia."


def _max_severity(values: list[str]) -> str:
    return max(values, key=lambda item: _SEVERITY_RANK.get(item, 0)) if values else "baja"


def _snippet(text: str, start: int, end: int, radius: int = 80) -> str:
    left = max(0, start - radius)
    right = min(len(text), end + radius)
    prefix = "..." if left > 0 else ""
    suffix = "..." if right < len(text) else ""
    return prefix + " ".join(text[left:right].split()) + suffix


def _redact_sensitive_text(text: str) -> str:
    redacted = text
    for rule in _DATA_PROTECTION_RULES:
        redacted = re.sub(rule["pattern"], f"[{rule['label'].upper()}]", redacted, flags=re.IGNORECASE)
    return redacted


def _redacted_preview(text: str) -> str:
    clean = " ".join((text or "").split())
    return _redact_sensitive_text(clean[:1200])


def list_contract_changes(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_changes WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["changes"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
        items.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("chg"), "items": items}


def create_contract_change(user: UserContext, contract_id: str, payload: ContractChangeCreate) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    change_id = _id("CHG")
    proposed_end_date = payload.proposed_end_date or _add_months(contract.get("end_date"), payload.extension_months)
    default_title = "Propuesta de prorroga" if payload.change_type == "extension" else "Propuesta de modificacion contractual"
    item = {
        "id": change_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "change_type": payload.change_type,
        "title": payload.title or default_title,
        "reason": payload.reason,
        "legal_basis": payload.legal_basis,
        "impact_summary": payload.impact_summary,
        "amount_delta": payload.amount_delta,
        "proposed_start_date": payload.proposed_start_date,
        "proposed_end_date": proposed_end_date,
        "extension_months": payload.extension_months,
        "source_refs": payload.source_refs,
        "document_id": None,
        "ai_draft": True,
        "status": "draft",
        "notes": payload.notes,
        "created_by": user.user_id,
        "validated_by": None,
        "validated_at": None,
        "applied_by": None,
        "applied_at": None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_changes(
                  id, tenant_id, contract_id, change_type, title, reason, legal_basis,
                  impact_summary, amount_delta, proposed_start_date, proposed_end_date,
                  extension_months, source_refs, document_id, ai_draft, status, notes, created_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NULL, TRUE, 'draft', %s, %s)
                """,
                (
                    change_id,
                    user.tenant_id,
                    contract_id,
                    payload.change_type,
                    item["title"],
                    payload.reason,
                    payload.legal_basis,
                    payload.impact_summary,
                    payload.amount_delta,
                    payload.proposed_start_date,
                    proposed_end_date,
                    payload.extension_months,
                    Json(payload.source_refs),
                    payload.notes,
                    user.user_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["changes"][change_id] = item
    _audit(user, "change.created", contract_id, {"change_id": change_id, "change_type": payload.change_type})
    return get_contract_change(user, change_id) or item


def get_contract_change(user: UserContext, change_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.contract_changes WHERE tenant_id = %s AND id = %s", (user.tenant_id, change_id)))
    item = _MEMORY["changes"].get(change_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def update_contract_change(user: UserContext, change_id: str, payload: ContractChangeUpdate) -> dict[str, Any]:
    current = get_contract_change(user, change_id)
    if not current:
        raise KeyError(change_id)
    updates = payload.model_dump(exclude_unset=True)
    merged = {**current, **updates}
    if "proposed_end_date" not in updates and updates.get("extension_months"):
        contract = get_contract(user, current["contract_id"])
        merged["proposed_end_date"] = _add_months(contract.get("end_date") if contract else None, updates.get("extension_months"))
    if updates.get("status") == "validated":
        merged["validated_by"] = user.user_id
        merged["validated_at"] = now_iso()
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contract_changes
                SET title = %s, reason = %s, legal_basis = %s, impact_summary = %s,
                    amount_delta = %s, proposed_start_date = %s, proposed_end_date = %s,
                    extension_months = %s, source_refs = %s, document_id = %s, status = %s,
                    notes = %s, validated_by = %s, validated_at = %s, updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (
                    merged.get("title"),
                    merged.get("reason"),
                    merged.get("legal_basis"),
                    merged.get("impact_summary"),
                    merged.get("amount_delta"),
                    merged.get("proposed_start_date"),
                    merged.get("proposed_end_date"),
                    merged.get("extension_months"),
                    Json(merged.get("source_refs") or []),
                    merged.get("document_id"),
                    merged.get("status"),
                    merged.get("notes"),
                    merged.get("validated_by"),
                    merged.get("validated_at"),
                    user.tenant_id,
                    change_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["changes"][change_id].update(merged)
        _MEMORY["changes"][change_id]["updated_at"] = now_iso()
    _audit(user, "change.updated", current["contract_id"], {"change_id": change_id, "fields": sorted(updates)})
    return get_contract_change(user, change_id) or merged


def generate_contract_change_report(user: UserContext, change_id: str) -> dict[str, Any]:
    change = get_contract_change(user, change_id)
    if not change:
        raise KeyError(change_id)
    document_type = "prorroga_informe_juridico" if change["change_type"] == "extension" else "modificacion_informe_juridico"
    notes = "\n".join(
        [
            f"Propuesta: {change.get('title')}",
            f"Motivo: {change.get('reason')}",
            f"Base juridica: {change.get('legal_basis') or 'pendiente de completar'}",
            f"Impacto: {change.get('impact_summary') or 'pendiente de completar'}",
            f"Variacion economica: {change.get('amount_delta') if change.get('amount_delta') is not None else 'sin variacion indicada'}",
            f"Nueva fecha propuesta: {change.get('proposed_end_date') or 'sin cambio de fecha indicado'}",
            "Este informe es un borrador y requiere validacion humana.",
        ]
    )
    document = generate_document(user, change["contract_id"], GeneratedDocumentRequest(document_type=document_type, language="es", notes=notes))["document"]
    updated = update_contract_change(user, change_id, ContractChangeUpdate(document_id=document["id"], status="under_review"))
    _audit(user, "change.report_generated", change["contract_id"], {"change_id": change_id, "document_id": document["id"]})
    return {"trace_id": trace_id("chg"), "change": updated, "document": document}


def apply_contract_change(user: UserContext, change_id: str) -> dict[str, Any]:
    change = get_contract_change(user, change_id)
    if not change:
        raise KeyError(change_id)
    if change.get("status") != "validated":
        raise ValueError("change_must_be_validated_before_apply")
    contract = get_contract(user, change["contract_id"])
    if not contract:
        raise KeyError(change["contract_id"])
    data = dict(contract.get("data") or {})
    applied_changes = list(data.get("applied_changes") or [])
    applied_changes.append({"id": change_id, "type": change.get("change_type"), "applied_at": now_iso(), "title": change.get("title")})
    data["applied_changes"] = applied_changes
    next_awarded = _amount(contract.get("awarded_amount")) + _amount(change.get("amount_delta")) if change.get("amount_delta") is not None else contract.get("awarded_amount")
    next_end_date = change.get("proposed_end_date") or contract.get("end_date")
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contracts
                SET awarded_amount = %s, end_date = %s, data = %s, updated_by = %s, updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (next_awarded, next_end_date, Json(data), user.user_id, user.tenant_id, change["contract_id"]),
            )
            cursor.execute(
                """
                UPDATE xfollow.contract_changes
                SET status = 'applied', applied_by = %s, applied_at = now(), updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (user.user_id, user.tenant_id, change_id),
            )
            conn.commit()
    else:
        _MEMORY["contracts"][change["contract_id"]]["awarded_amount"] = next_awarded
        _MEMORY["contracts"][change["contract_id"]]["end_date"] = next_end_date
        _MEMORY["contracts"][change["contract_id"]]["data"] = data
        _MEMORY["contracts"][change["contract_id"]]["updated_by"] = user.user_id
        _MEMORY["contracts"][change["contract_id"]]["updated_at"] = now_iso()
        _MEMORY["changes"][change_id]["status"] = "applied"
        _MEMORY["changes"][change_id]["applied_by"] = user.user_id
        _MEMORY["changes"][change_id]["applied_at"] = now_iso()
        _MEMORY["changes"][change_id]["updated_at"] = now_iso()
    _audit(user, "change.applied", change["contract_id"], {"change_id": change_id, "end_date": next_end_date, "awarded_amount": next_awarded})
    return {"trace_id": trace_id("chg"), "change": get_contract_change(user, change_id), "contract": get_contract(user, change["contract_id"])}


def list_invoices(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_invoices WHERE tenant_id = %s AND contract_id = %s ORDER BY invoice_date DESC NULLS LAST, created_at DESC",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["invoices"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
        items.sort(key=lambda item: item.get("invoice_date") or item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("inv"), "items": items}


def create_invoice(user: UserContext, contract_id: str, payload: InvoiceCreate) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    invoice_id = _id("INV")
    amount_with_tax = payload.amount_with_tax if payload.amount_with_tax is not None else payload.amount_without_tax + payload.tax_amount
    payment_due_date = payload.payment_due_date or _add_days(payload.invoice_date, 30)
    item = {
        "id": invoice_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "status": "registered",
        "compliance_check_id": None,
        "paid_at": None,
        "registered_by": user.user_id,
        "validated_by": None,
        "validated_at": None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
        **payload.model_dump(),
        "supplier": payload.supplier or contract.get("contractor"),
        "amount_with_tax": round(amount_with_tax, 2),
        "payment_due_date": payment_due_date,
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_invoices(
                  id, tenant_id, contract_id, obligation_id, invoice_number, supplier, concept,
                  invoice_date, service_period_start, service_period_end, amount_without_tax,
                  tax_amount, amount_with_tax, currency, status, evidence_ids, payment_due_date,
                  notes, registered_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'registered', %s, %s, %s, %s)
                """,
                (
                    invoice_id,
                    user.tenant_id,
                    contract_id,
                    payload.obligation_id,
                    payload.invoice_number,
                    payload.supplier or contract.get("contractor"),
                    payload.concept,
                    payload.invoice_date,
                    payload.service_period_start,
                    payload.service_period_end,
                    payload.amount_without_tax,
                    payload.tax_amount,
                    round(amount_with_tax, 2),
                    payload.currency,
                    Json(payload.evidence_ids),
                    payment_due_date,
                    payload.notes,
                    user.user_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["invoices"][invoice_id] = item
    _ensure_invoice_alert(user, contract_id, invoice_id, payload.invoice_number, payment_due_date)
    _audit(user, "invoice.created", contract_id, {"invoice_id": invoice_id, "invoice_number": payload.invoice_number, "amount_with_tax": round(amount_with_tax, 2)})
    return get_invoice(user, invoice_id) or item


def get_invoice(user: UserContext, invoice_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.contract_invoices WHERE tenant_id = %s AND id = %s", (user.tenant_id, invoice_id)))
    item = _MEMORY["invoices"].get(invoice_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def update_invoice(user: UserContext, invoice_id: str, payload: InvoiceUpdate) -> dict[str, Any]:
    current = get_invoice(user, invoice_id)
    if not current:
        raise KeyError(invoice_id)
    updates = payload.model_dump(exclude_unset=True)
    merged = {**current, **updates}
    if "amount_with_tax" not in updates and ({"amount_without_tax", "tax_amount"} & set(updates)):
        merged["amount_with_tax"] = round(_amount(merged.get("amount_without_tax")) + _amount(merged.get("tax_amount")), 2)
    if not merged.get("payment_due_date"):
        merged["payment_due_date"] = _add_days(merged.get("invoice_date"), 30)
    if updates.get("status") in {"conforming", "payment_ordered", "paid"}:
        merged["validated_by"] = merged.get("validated_by") or user.user_id
        merged["validated_at"] = merged.get("validated_at") or now_iso()
    if updates.get("status") == "paid" and not merged.get("paid_at"):
        merged["paid_at"] = now_iso()
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contract_invoices
                SET obligation_id = %s, invoice_number = %s, supplier = %s, concept = %s,
                    invoice_date = %s, service_period_start = %s, service_period_end = %s,
                    amount_without_tax = %s, tax_amount = %s, amount_with_tax = %s,
                    currency = %s, status = %s, evidence_ids = %s, payment_due_date = %s,
                    paid_at = %s, notes = %s, validated_by = %s, validated_at = %s,
                    updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (
                    merged.get("obligation_id"),
                    merged.get("invoice_number"),
                    merged.get("supplier"),
                    merged.get("concept"),
                    merged.get("invoice_date"),
                    merged.get("service_period_start"),
                    merged.get("service_period_end"),
                    merged.get("amount_without_tax"),
                    merged.get("tax_amount"),
                    merged.get("amount_with_tax"),
                    merged.get("currency") or "EUR",
                    merged.get("status"),
                    Json(merged.get("evidence_ids") or []),
                    merged.get("payment_due_date"),
                    merged.get("paid_at"),
                    merged.get("notes"),
                    merged.get("validated_by"),
                    merged.get("validated_at"),
                    user.tenant_id,
                    invoice_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["invoices"][invoice_id].update(merged)
        _MEMORY["invoices"][invoice_id]["updated_at"] = now_iso()
    if merged.get("status") not in {"paid", "rejected"}:
        _ensure_invoice_alert(user, merged["contract_id"], invoice_id, merged.get("invoice_number") or invoice_id, merged.get("payment_due_date"))
    _audit(user, "invoice.updated", current["contract_id"], {"invoice_id": invoice_id, "fields": sorted(updates)})
    return get_invoice(user, invoice_id) or merged


def _ensure_invoice_alert(user: UserContext, contract_id: str, invoice_id: str, invoice_number: str, due_date: str | None) -> None:
    if not due_date:
        return
    title = f"Factura {invoice_number}: vencimiento pago"
    if db.db_available():
        existing = db.fetch_one(
            "SELECT id FROM xfollow.contract_alerts WHERE tenant_id = %s AND contract_id = %s AND title = %s",
            (user.tenant_id, contract_id, title),
        )
        if existing:
            return
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_alerts(id, tenant_id, contract_id, obligation_id, alert_type, title, due_date, lead_days, severity)
                VALUES (%s, %s, %s, NULL, 'payment', %s, %s, 7, 'alta')
                """,
                (_id("ALT"), user.tenant_id, contract_id, title, due_date),
            )
            conn.commit()
    else:
        if any(item["contract_id"] == contract_id and item["title"] == title for item in _MEMORY["alerts"].values()):
            return
        alert_id = _id("ALT")
        _MEMORY["alerts"][alert_id] = {
            "id": alert_id,
            "tenant_id": user.tenant_id,
            "contract_id": contract_id,
            "obligation_id": None,
            "alert_type": "payment",
            "title": title,
            "due_date": due_date,
            "lead_days": 7,
            "severity": "alta",
            "status": "open",
            "generated_by": "system",
            "created_at": now_iso(),
            "updated_at": now_iso(),
            "source_invoice_id": invoice_id,
        }


def liquidation_preview(user: UserContext, contract_id: str) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    obligations = list_obligations(user, contract_id)["items"]
    checks = list_checks(user, contract_id)["items"]
    invoices = list_invoices(user, contract_id)["items"]
    penalties = list_penalties(user, contract_id)["items"]
    proceedings = list_proceedings(user, contract_id)["items"]
    changes = list_contract_changes(user, contract_id)["items"]
    return {
        "trace_id": trace_id("liq"),
        "summary": _liquidation_summary(contract, invoices, penalties, obligations, checks, proceedings, changes),
        "invoices": invoices,
        "penalties": penalties,
    }


def _liquidation_summary(
    contract: dict[str, Any],
    invoices: list[dict[str, Any]],
    penalties: list[dict[str, Any]],
    obligations: list[dict[str, Any]],
    checks: list[dict[str, Any]],
    proceedings: list[dict[str, Any]],
    changes: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    changes = changes or []
    accepted_statuses = {"conforming", "payment_ordered", "paid"}
    awarded_amount = _amount(contract.get("awarded_amount") or contract.get("budget_with_tax") or contract.get("budget_without_tax"))
    invoiced_amount = sum(_amount(item.get("amount_with_tax")) for item in invoices if item.get("status") != "rejected")
    conforming_amount = sum(_amount(item.get("amount_with_tax")) for item in invoices if item.get("status") in accepted_statuses)
    paid_amount = sum(_amount(item.get("amount_with_tax")) for item in invoices if item.get("status") == "paid")
    pending_payment_amount = sum(_amount(item.get("amount_with_tax")) for item in invoices if item.get("status") in {"conforming", "payment_ordered"})
    penalties_amount = sum(_amount(item.get("calculated_amount")) for item in penalties if item.get("status") in {"validated", "opened", "resolved"})
    draft_obligations = sum(1 for item in obligations if item.get("validation_status") == "draft")
    risky_checks = sum(1 for item in checks if item.get("result") in {"at_risk", "breach", "needs_review"} and item.get("status") != "rejected")
    open_proceedings = sum(1 for item in proceedings if item.get("status") not in {"resolved", "closed", "rejected"})
    open_changes = sum(1 for item in changes if item.get("status") not in {"applied", "rejected"})
    settlement_balance = round(conforming_amount - paid_amount - penalties_amount, 2)
    remaining_authorized_budget = round(awarded_amount - conforming_amount, 2) if awarded_amount else None
    warnings = []
    if draft_obligations:
        warnings.append(f"{draft_obligations} obligaciones pendientes de validacion")
    if risky_checks:
        warnings.append(f"{risky_checks} verificaciones requieren revision")
    if pending_payment_amount:
        warnings.append("Hay facturas conformadas pendientes de pago")
    if open_proceedings:
        warnings.append("Hay expedientes abiertos antes del cierre")
    if open_changes:
        warnings.append("Hay propuestas de modificacion o prorroga sin aplicar")
    return {
        "awarded_amount": round(awarded_amount, 2),
        "invoiced_amount": round(invoiced_amount, 2),
        "conforming_amount": round(conforming_amount, 2),
        "paid_amount": round(paid_amount, 2),
        "pending_payment_amount": round(pending_payment_amount, 2),
        "validated_penalties_amount": round(penalties_amount, 2),
        "settlement_balance": settlement_balance,
        "remaining_authorized_budget": remaining_authorized_budget,
        "draft_obligations": draft_obligations,
        "risky_checks": risky_checks,
        "open_proceedings": open_proceedings,
        "open_changes": open_changes,
        "warnings": warnings,
        "requires_human_validation": True,
    }


def list_proceedings(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.contract_proceedings WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["proceedings"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
        items.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("prc"), "items": items}


def open_proceeding(user: UserContext, contract_id: str, payload: ProceedingCreate) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    penalty = _get_penalty(user, payload.penalty_case_id) if payload.penalty_case_id else None
    if penalty and penalty.get("contract_id") != contract_id:
        raise KeyError(payload.penalty_case_id or "penalty_case")
    proceeding_id = _id("PRC")
    title = payload.title or ("Expediente de resolucion contractual" if payload.case_type == "resolution" else "Expediente de imposicion de penalidad")
    basis = payload.legal_basis or (penalty.get("basis_text") if penalty else None)
    summary = payload.trigger_text
    if penalty:
        summary = f"{payload.trigger_text}\n\nCaso vinculado: {penalty['id']} por {penalty.get('calculated_amount')} EUR."
    timeline = [{"at": now_iso(), "status": "opened", "user_id": user.user_id, "comment": "Apertura propuesta. Requiere tramitacion y validacion humana."}]
    item = {
        "id": proceeding_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "penalty_case_id": payload.penalty_case_id,
        "case_type": payload.case_type,
        "title": title,
        "summary": summary,
        "legal_basis": basis,
        "proposed_action": payload.proposed_action,
        "status": "opened",
        "due_date": payload.due_date,
        "document_ids": payload.document_ids,
        "timeline": timeline,
        "created_by": user.user_id,
        "updated_by": user.user_id,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.contract_proceedings(
                  id, tenant_id, contract_id, penalty_case_id, case_type, title, summary,
                  legal_basis, proposed_action, status, due_date, document_ids, timeline,
                  created_by, updated_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'opened', %s, %s, %s, %s, %s)
                """,
                (
                    proceeding_id,
                    user.tenant_id,
                    contract_id,
                    payload.penalty_case_id,
                    payload.case_type,
                    title,
                    summary,
                    basis,
                    payload.proposed_action,
                    payload.due_date,
                    Json(payload.document_ids),
                    Json(timeline),
                    user.user_id,
                    user.user_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["proceedings"][proceeding_id] = item
    _audit(user, "proceeding.opened", contract_id, {"proceeding_id": proceeding_id, "case_type": payload.case_type, "penalty_case_id": payload.penalty_case_id})
    return get_proceeding(user, proceeding_id) or item


def get_proceeding(user: UserContext, proceeding_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.contract_proceedings WHERE tenant_id = %s AND id = %s", (user.tenant_id, proceeding_id)))
    item = _MEMORY["proceedings"].get(proceeding_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def update_proceeding(user: UserContext, proceeding_id: str, payload: ProceedingUpdate) -> dict[str, Any]:
    current = get_proceeding(user, proceeding_id)
    if not current:
        raise KeyError(proceeding_id)
    updates = payload.model_dump(exclude_unset=True)
    merged = {**current, **{key: value for key, value in updates.items() if key != "comment"}}
    timeline = list(merged.get("timeline") or [])
    if updates.get("status") or updates.get("comment"):
        timeline.append({"at": now_iso(), "status": merged.get("status"), "user_id": user.user_id, "comment": updates.get("comment")})
    merged["timeline"] = timeline
    if updates.get("status") in {"resolved", "closed"}:
        merged["validated_by"] = user.user_id
        merged["validated_at"] = now_iso()
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.contract_proceedings
                SET title = %s, summary = %s, legal_basis = %s, proposed_action = %s,
                    status = %s, due_date = %s, document_ids = %s, timeline = %s,
                    updated_by = %s, validated_by = %s, validated_at = %s, updated_at = now()
                WHERE tenant_id = %s AND id = %s
                """,
                (
                    merged.get("title"),
                    merged.get("summary"),
                    merged.get("legal_basis"),
                    merged.get("proposed_action"),
                    merged.get("status"),
                    merged.get("due_date"),
                    Json(merged.get("document_ids") or []),
                    Json(timeline),
                    user.user_id,
                    merged.get("validated_by"),
                    merged.get("validated_at"),
                    user.tenant_id,
                    proceeding_id,
                ),
            )
            conn.commit()
    else:
        _MEMORY["proceedings"][proceeding_id].update(merged)
        _MEMORY["proceedings"][proceeding_id]["updated_by"] = user.user_id
        _MEMORY["proceedings"][proceeding_id]["updated_at"] = now_iso()
    _audit(user, "proceeding.updated", current["contract_id"], {"proceeding_id": proceeding_id, "fields": sorted(updates)})
    return get_proceeding(user, proceeding_id) or merged


def run_compliance_checks(user: UserContext, contract_id: str, request: ComplianceCheckRequest) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    obligations = list_obligations(user, contract_id)["items"]
    if request.obligation_id:
        obligations = [item for item in obligations if item["id"] == request.obligation_id]
    else:
        obligations = [item for item in obligations if item.get("validation_status") == "validated"] or obligations
    if not obligations:
        raise KeyError("obligation_not_found")
    evidence = list_evidence(user, contract_id)["items"]
    if request.evidence_ids:
        evidence = [item for item in evidence if item["id"] in set(request.evidence_ids)]
    created = []
    for obligation in obligations:
        related = [item for item in evidence if not item.get("obligation_id") or item.get("obligation_id") == obligation["id"]]
        combined = "\n".join(str(item.get("extracted_text") or item.get("title") or "") for item in related).lower()
        if not related:
            result = "needs_review"
            reasoning = "No hay evidencias asociadas suficientes para verificar la obligacion."
        elif "incumpl" in combined or "retras" in combined:
            result = "breach"
            reasoning = "La evidencia contiene indicios de incumplimiento o retraso; requiere revision humana."
        elif "riesgo" in combined or "pendiente" in combined:
            result = "at_risk"
            reasoning = "La evidencia sugiere riesgo o tareas pendientes."
        elif "cumpl" in combined or "entreg" in combined or "recibid" in combined:
            result = "compliant"
            reasoning = "La evidencia aportada parece respaldar el cumplimiento, pendiente de validacion humana."
        else:
            result = "needs_review"
            reasoning = "La evidencia no permite una conclusion automatica fiable."
        created.append(_insert_check(user, contract_id, obligation["id"], [item["id"] for item in related], result, reasoning, related))
    _audit(user, "compliance.checked", contract_id, {"checks": len(created)})
    return {"trace_id": trace_id("chk"), "items": created}


def _insert_check(user: UserContext, contract_id: str, obligation_id: str, evidence_ids: list[str], result: str, reasoning: str, evidence: list[dict[str, Any]]) -> dict[str, Any]:
    check_id = _id("CHK")
    citations = [{"evidence_id": item["id"], "title": item["title"], "source_url": item.get("source_url")} for item in evidence]
    item = {
        "id": check_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "obligation_id": obligation_id,
        "evidence_ids": evidence_ids,
        "result": result,
        "reasoning": reasoning,
        "citations": citations,
        "ai_draft": True,
        "status": "draft",
        "created_by": user.user_id,
        "created_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.compliance_checks(id, tenant_id, contract_id, obligation_id, evidence_ids, result, reasoning, citations, created_by)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (check_id, user.tenant_id, contract_id, obligation_id, Json(evidence_ids), result, reasoning, Json(citations), user.user_id),
            )
            conn.commit()
    else:
        _MEMORY["checks"][check_id] = item
    return _row(db.fetch_one("SELECT * FROM xfollow.compliance_checks WHERE tenant_id = %s AND id = %s", (user.tenant_id, check_id))) if db.db_available() else item


def create_ai_obligations(user: UserContext, contract_id: str, result: Any, sources: list[dict[str, Any]]) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    raw_items = result.get("items") if isinstance(result, dict) else None
    if not isinstance(raw_items, list):
        raise ValueError("ai_obligations_invalid_result")
    source_by_id = {str(item["id"]): item for item in sources}
    existing_titles = {str(item.get("title") or "").lower() for item in list_obligations(user, contract_id)["items"]}
    allowed_categories = {"deadline", "delivery", "sla", "special_condition", "warranty", "payment", "documentation", "extension", "other"}
    allowed_severity = {"baja", "media", "alta", "critica"}
    created: list[dict[str, Any]] = []
    for raw in raw_items[:24]:
        if not isinstance(raw, dict):
            continue
        title = str(raw.get("title") or "").strip()[:240]
        description = str(raw.get("description") or "").strip()[:5000]
        if not title or not description or title.lower() in existing_titles:
            continue
        category = str(raw.get("category") or "other")
        severity = str(raw.get("severity") or "media")
        if category not in allowed_categories:
            category = "other"
        if severity not in allowed_severity:
            severity = "media"
        source_ids = [str(value) for value in raw.get("source_ids") or [] if str(value) in source_by_id]
        source_refs = [
            {"source_kind": "contract_evidence", "source_id": source_id, "title": source_by_id[source_id].get("title")}
            for source_id in source_ids
        ]
        proposal = ObligationCreate(
            category=category,  # type: ignore[arg-type]
            title=title,
            description=description,
            due_date=raw.get("due_date") or None,
            severity=severity,  # type: ignore[arg-type]
            source_refs=source_refs,
        )
        created.append(create_obligation(user, contract_id, proposal))
        existing_titles.add(title.lower())
    _audit(user, "obligations.ai_extracted", contract_id, {"created": len(created)})
    return {"items": created}


def create_ai_checks(user: UserContext, contract_id: str, result: Any, sources: list[dict[str, Any]]) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    raw_items = result.get("items") if isinstance(result, dict) else None
    if not isinstance(raw_items, list):
        raise ValueError("ai_checks_invalid_result")
    source_by_id = {str(item["id"]): item for item in sources}
    obligations = {item["id"]: item for item in list_obligations(user, contract_id)["items"]}
    allowed_results = {"compliant", "at_risk", "breach", "needs_review"}
    created: list[dict[str, Any]] = []
    for raw in raw_items[:24]:
        if not isinstance(raw, dict):
            continue
        obligation_id = str(raw.get("obligation_id") or "")
        if obligation_id not in obligations:
            continue
        check_result = str(raw.get("result") or "needs_review")
        if check_result not in allowed_results:
            check_result = "needs_review"
        reasoning = str(raw.get("reasoning") or "La IA no ha aportado una conclusion verificable.").strip()[:5000]
        evidence_ids = [str(value) for value in raw.get("evidence_ids") or [] if str(value) in source_by_id]
        evidence = [source_by_id[value] for value in evidence_ids]
        created.append(_insert_check(user, contract_id, obligation_id, evidence_ids, check_result, reasoning, evidence))
    _audit(user, "compliance.ai_checked", contract_id, {"checks": len(created)})
    return {"items": created}


def list_checks(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all("SELECT * FROM xfollow.compliance_checks WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC", (user.tenant_id, contract_id))
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["checks"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("chk"), "items": items}


def resolve_check(user: UserContext, check_id: str, request: ResolutionRequest) -> dict[str, Any]:
    check = _get_check(user, check_id)
    if not check:
        raise KeyError(check_id)
    if request.status not in {"validated", "rejected"}:
        raise ValueError("invalid_check_status")
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.compliance_checks
                SET status = %s,
                    validated_by = CASE WHEN %s = 'validated' THEN %s ELSE validated_by END,
                    validated_at = CASE WHEN %s = 'validated' THEN now() ELSE validated_at END
                WHERE tenant_id = %s AND id = %s
                """,
                (request.status, request.status, user.user_id, request.status, user.tenant_id, check_id),
            )
            conn.commit()
    else:
        _MEMORY["checks"][check_id]["status"] = request.status
        if request.status == "validated":
            _MEMORY["checks"][check_id]["validated_by"] = user.user_id
            _MEMORY["checks"][check_id]["validated_at"] = now_iso()
    _audit(user, "compliance.resolved", check["contract_id"], {"check_id": check_id, "status": request.status, "comment": request.comment})
    return _get_check(user, check_id) or check


def _get_check(user: UserContext, check_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.compliance_checks WHERE tenant_id = %s AND id = %s", (user.tenant_id, check_id)))
    item = _MEMORY["checks"].get(check_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def generate_document(user: UserContext, contract_id: str, request: GeneratedDocumentRequest) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    obligations = list_obligations(user, contract_id)["items"]
    checks = list_checks(user, contract_id)["items"]
    penalties = list_penalties(user, contract_id)["items"]
    invoices = list_invoices(user, contract_id)["items"]
    proceedings = list_proceedings(user, contract_id)["items"]
    changes = list_contract_changes(user, contract_id)["items"]
    titles = {
        "recepcion_acta": "Acta de recepcion",
        "liquidacion_informe": "Informe de liquidacion",
        "modificacion_informe_juridico": "Informe juridico de modificacion",
        "prorroga_informe_juridico": "Informe juridico de prorroga",
        "penalidad_expediente": "Expediente de imposicion de penalidad",
        "resolucion_expediente": "Expediente de resolucion contractual",
    }
    title = titles[request.document_type]
    content = _document_content(title, contract, obligations, checks, penalties, invoices, proceedings, changes, request.notes)
    citations = [{"type": "contract", "id": contract_id, "title": contract["title"]}]
    item = _store_generated_document(
        user,
        contract_id,
        request.document_type,
        title,
        content,
        citations,
        {"mode": "deterministic_draft", "language": request.language},
    )

    _audit(user, "document.generated", contract_id, {"document_id": item["id"], "document_type": request.document_type})
    return {"trace_id": trace_id("doc"), "document": item}


def _store_generated_document(
    user: UserContext,
    contract_id: str,
    document_type: str,
    title: str,
    content: str,
    citations: list[dict[str, Any]],
    prompt_trace: dict[str, Any],
) -> dict[str, Any]:
    document_id = _id("DOC")
    item = {
        "id": document_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "document_type": document_type,
        "title": title,
        "content_text": content,
        "citations": citations,
        "prompt_trace": prompt_trace,
        "status": "draft",
        "generated_by": user.user_id,
        "created_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.generated_documents(id, tenant_id, contract_id, document_type, title, content_text, citations, prompt_trace, generated_by)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (document_id, user.tenant_id, contract_id, document_type, title, content, Json(citations), Json(prompt_trace), user.user_id),
            )
            conn.commit()
    else:
        _MEMORY["documents"][document_id] = item
    return _get_document(user, document_id) or item


def create_ai_document(
    user: UserContext,
    contract_id: str,
    request: GeneratedDocumentRequest,
    content: str,
    sources: list[dict[str, Any]],
) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    titles = {
        "recepcion_acta": "Acta de recepcion",
        "liquidacion_informe": "Informe de liquidacion",
        "modificacion_informe_juridico": "Informe juridico de modificacion",
        "prorroga_informe_juridico": "Informe juridico de prorroga",
        "penalidad_expediente": "Expediente de imposicion de penalidad",
        "resolucion_expediente": "Expediente de resolucion contractual",
    }
    citations = [
        {"type": "contract_evidence", "id": source["id"], "title": source.get("title")}
        for source in sources
    ]
    document = _store_generated_document(
        user,
        contract_id,
        request.document_type,
        titles[request.document_type],
        content.strip(),
        citations,
        {
            "mode": "llm2",
            "model": load_runtime_config().llm_model,
            "language": request.language,
            "source_ids": [source["id"] for source in sources],
        },
    )
    _audit(user, "document.ai_generated", contract_id, {"document_id": document["id"], "document_type": request.document_type})
    return document


def _document_content(
    title: str,
    contract: dict[str, Any],
    obligations: list[dict[str, Any]],
    checks: list[dict[str, Any]],
    penalties: list[dict[str, Any]],
    invoices: list[dict[str, Any]],
    proceedings: list[dict[str, Any]],
    changes: list[dict[str, Any]],
    notes: str | None,
) -> str:
    open_obligations = [item for item in obligations if item.get("status") not in {"met", "waived"}]
    risky_checks = [item for item in checks if item.get("result") in {"at_risk", "breach", "needs_review"}]
    liquidation = _liquidation_summary(contract, invoices, penalties, obligations, checks, proceedings, changes)
    lines = [
        f"# {title}",
        "",
        "Borrador generado por xFollow. Requiere revision y validacion humana antes de cualquier efecto administrativo.",
        "",
        "## Contrato",
        f"- Expediente: {contract.get('file_number') or contract['id']}",
        f"- Objeto: {contract['title']}",
        f"- Organo de contratacion: {contract.get('contracting_body') or 'pendiente'}",
        f"- Contratista: {contract.get('contractor') or 'pendiente'}",
        f"- Periodo: {contract.get('start_date') or 'pendiente'} - {contract.get('end_date') or 'pendiente'}",
        "",
        "## Situacion de ejecucion",
        f"- Obligaciones registradas: {len(obligations)}",
        f"- Obligaciones abiertas: {len(open_obligations)}",
        f"- Verificaciones con riesgo o pendientes: {len(risky_checks)}",
        f"- Expedientes de penalidad/resolucion: {len(penalties)}",
        "",
        "## Situacion economica",
        f"- Importe adjudicado/base de control: {liquidation['awarded_amount']} EUR",
        f"- Facturado: {liquidation['invoiced_amount']} EUR",
        f"- Conformado: {liquidation['conforming_amount']} EUR",
        f"- Pagado: {liquidation['paid_amount']} EUR",
        f"- Pendiente de pago: {liquidation['pending_payment_amount']} EUR",
        f"- Penalidades validadas/abiertas: {liquidation['validated_penalties_amount']} EUR",
        f"- Saldo de liquidacion propuesto: {liquidation['settlement_balance']} EUR",
    ]
    if invoices:
        lines.extend(["", "## Facturas registradas"])
        lines.extend(f"- {item.get('invoice_number')}: {item.get('status')} · {item.get('amount_with_tax')} {item.get('currency')}" for item in invoices[:10])
    if risky_checks:
        lines.extend(["", "## Incidencias a revisar"])
        lines.extend(f"- {item.get('result')}: {item.get('reasoning')}" for item in risky_checks[:8])
    open_proceedings = [item for item in proceedings if item.get("status") not in {"resolved", "closed", "rejected"}]
    if open_proceedings:
        lines.extend(["", "## Expedientes abiertos"])
        lines.extend(f"- {item.get('title')}: {item.get('status')}" for item in open_proceedings[:8])
    open_changes = [item for item in changes if item.get("status") not in {"applied", "rejected"}]
    if open_changes:
        lines.extend(["", "## Modificaciones y prorrogas en curso"])
        lines.extend(f"- {item.get('title')}: {item.get('status')} · {item.get('proposed_end_date') or 'sin fecha propuesta'}" for item in open_changes[:8])
    if liquidation["warnings"]:
        lines.extend(["", "## Advertencias de cierre"])
        lines.extend(f"- {warning}" for warning in liquidation["warnings"])
    if notes:
        lines.extend(["", "## Observaciones del usuario", notes])
    lines.extend(["", "## Validacion", "Este borrador debe ser revisado por la unidad responsable y, cuando proceda, por servicios juridicos/intervencion."])
    return "\n".join(lines)


def _get_document(user: UserContext, document_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.generated_documents WHERE tenant_id = %s AND id = %s", (user.tenant_id, document_id)))
    item = _MEMORY["documents"].get(document_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def get_document(user: UserContext, document_id: str) -> dict[str, Any] | None:
    return _get_document(user, document_id)


def resolve_document(user: UserContext, document_id: str, request: ResolutionRequest) -> dict[str, Any]:
    document = _get_document(user, document_id)
    if not document:
        raise KeyError(document_id)
    if request.status not in {"validated", "rejected"}:
        raise ValueError("invalid_document_status")
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.generated_documents
                SET status = %s,
                    validated_by = CASE WHEN %s = 'validated' THEN %s ELSE validated_by END,
                    validated_at = CASE WHEN %s = 'validated' THEN now() ELSE validated_at END
                WHERE tenant_id = %s AND id = %s
                """,
                (request.status, request.status, user.user_id, request.status, user.tenant_id, document_id),
            )
            conn.commit()
    else:
        _MEMORY["documents"][document_id]["status"] = request.status
        if request.status == "validated":
            _MEMORY["documents"][document_id]["validated_by"] = user.user_id
            _MEMORY["documents"][document_id]["validated_at"] = now_iso()
    _audit(user, "document.resolved", document["contract_id"], {"document_id": document_id, "status": request.status, "comment": request.comment})
    return _get_document(user, document_id) or document


def list_documents(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all("SELECT * FROM xfollow.generated_documents WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC", (user.tenant_id, contract_id))
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["documents"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("doc"), "items": items}


def calculate_penalty(user: UserContext, contract_id: str, request: PenaltyCalculationRequest) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if request.daily_amount is not None and request.days_late is not None:
        raw = request.daily_amount * request.days_late
        method = "daily_amount_x_days_late"
    elif request.base_amount is not None and request.percentage is not None:
        raw = request.base_amount * request.percentage / 100
        method = "base_amount_x_percentage"
    else:
        raw = 0.0
        method = "manual_review_required"
    amount = min(raw, request.cap_amount) if request.cap_amount is not None else raw
    penalty_id = _id("PEN")
    calculation = {
        "method": method,
        "raw_amount": round(raw, 2),
        "cap_amount": request.cap_amount,
        "calculated_amount": round(amount, 2),
        "requires_human_validation": True,
    }
    content = (
        "# Propuesta de penalidad\n\n"
        f"Base: {request.basis_text}\n\n"
        f"Metodo: {method}. Importe calculado: {round(amount, 2)} EUR.\n\n"
        "La propuesta no tiene efecto ejecutivo sin tramitacion y validacion humana."
    )
    item = {
        "id": penalty_id,
        "tenant_id": user.tenant_id,
        "contract_id": contract_id,
        "obligation_id": request.obligation_id,
        "case_type": request.case_type,
        "basis_text": request.basis_text,
        "daily_amount": request.daily_amount,
        "days_late": request.days_late,
        "base_amount": request.base_amount,
        "percentage": request.percentage,
        "cap_amount": request.cap_amount,
        "calculated_amount": round(amount, 2),
        "calculation": calculation,
        "content_text": content,
        "citations": [],
        "status": "draft",
        "created_by": user.user_id,
        "created_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.penalty_cases(
                  id, tenant_id, contract_id, obligation_id, case_type, basis_text, daily_amount,
                  days_late, base_amount, percentage, cap_amount, calculated_amount, calculation,
                  content_text, citations, created_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (penalty_id, user.tenant_id, contract_id, request.obligation_id, request.case_type, request.basis_text, request.daily_amount, request.days_late, request.base_amount, request.percentage, request.cap_amount, round(amount, 2), Json(calculation), content, Json([]), user.user_id),
            )
            conn.commit()
    else:
        _MEMORY["penalties"][penalty_id] = item
    _audit(user, "penalty.calculated", contract_id, {"penalty_id": penalty_id, "calculated_amount": round(amount, 2)})
    return {"trace_id": trace_id("pen"), "penalty_case": _get_penalty(user, penalty_id) or item}


def _get_penalty(user: UserContext, penalty_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.penalty_cases WHERE tenant_id = %s AND id = %s", (user.tenant_id, penalty_id)))
    item = _MEMORY["penalties"].get(penalty_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def list_penalties(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all("SELECT * FROM xfollow.penalty_cases WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC", (user.tenant_id, contract_id))
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["penalties"].values() if item["tenant_id"] == user.tenant_id and item["contract_id"] == contract_id]
    return {"trace_id": trace_id("pen"), "items": items}


def resolve_penalty(user: UserContext, penalty_id: str, request: ResolutionRequest) -> dict[str, Any]:
    penalty = _get_penalty(user, penalty_id)
    if not penalty:
        raise KeyError(penalty_id)
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE xfollow.penalty_cases
                SET status = %s,
                    validated_by = CASE WHEN %s = 'validated' THEN %s ELSE validated_by END,
                    validated_at = CASE WHEN %s = 'validated' THEN now() ELSE validated_at END
                WHERE tenant_id = %s AND id = %s
                """,
                (request.status, request.status, user.user_id, request.status, user.tenant_id, penalty_id),
            )
            conn.commit()
    else:
        _MEMORY["penalties"][penalty_id]["status"] = request.status
        if request.status == "validated":
            _MEMORY["penalties"][penalty_id]["validated_by"] = user.user_id
            _MEMORY["penalties"][penalty_id]["validated_at"] = now_iso()
    _audit(user, "penalty.resolved", penalty["contract_id"], {"penalty_id": penalty_id, "status": request.status, "comment": request.comment})
    return _get_penalty(user, penalty_id) or penalty


def create_feedback(user: UserContext, payload: FeedbackCreate) -> dict[str, Any]:
    if payload.contract_id and not get_contract(user, payload.contract_id):
        raise KeyError(payload.contract_id)
    feedback_id = trace_id("fbk")
    item = {
        "id": feedback_id,
        "tenant_id": user.tenant_id,
        "contract_id": payload.contract_id,
        "artifact_type": payload.artifact_type,
        "artifact_id": payload.artifact_id,
        "rating": payload.rating,
        "comment": payload.comment,
        "metadata": payload.metadata,
        "created_by": user.user_id,
        "created_at": now_iso(),
    }
    if db.db_available():
        with db.connection() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO xfollow.ai_feedback(
                  id, tenant_id, contract_id, artifact_type, artifact_id, rating,
                  comment, metadata, created_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (feedback_id, user.tenant_id, payload.contract_id, payload.artifact_type, payload.artifact_id, payload.rating, payload.comment, Json(payload.metadata), user.user_id),
            )
            conn.commit()
    else:
        _MEMORY["feedback"][feedback_id] = item
    _audit(user, "feedback.created", payload.contract_id, {"feedback_id": feedback_id, "artifact_type": payload.artifact_type, "artifact_id": payload.artifact_id, "rating": payload.rating})
    return _get_feedback(user, feedback_id) or item


def _get_feedback(user: UserContext, feedback_id: str) -> dict[str, Any] | None:
    if db.db_available():
        return _row(db.fetch_one("SELECT * FROM xfollow.ai_feedback WHERE tenant_id = %s AND id = %s", (user.tenant_id, feedback_id)))
    item = _MEMORY["feedback"].get(feedback_id)
    return item if item and item["tenant_id"] == user.tenant_id else None


def list_feedback(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all("SELECT * FROM xfollow.ai_feedback WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC", (user.tenant_id, contract_id))
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["feedback"].values() if item["tenant_id"] == user.tenant_id and item.get("contract_id") == contract_id]
        items.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("fbk"), "items": items}


def governance_manifest(user: UserContext) -> dict[str, Any]:
    config = load_runtime_config()
    return {
        "trace_id": trace_id("gov"),
        "generated_at": now_iso(),
        "runtime": sanitized_config(config),
        "ai": {
            "user_disclosure": "La interfaz informa de que la persona usuaria interactua con asistencia de IA.",
            "generated_content_label": "Los documentos y verificaciones se marcan como borradores revisables.",
            "model": config.llm_model,
            "provider_mode": config.llm_base_url,
            "temperature": config.llm_temperature,
            "context_window": config.llm_context_window,
            "limitations": [
                "No produce actos administrativos por si misma.",
                "Las conclusiones deben contrastarse con expediente, evidencias y criterio profesional.",
                "Las reglas deterministas de demo deben sustituirse por conectores/modelos reales en despliegue final.",
            ],
        },
        "human_oversight": {
            "default_state": "draft",
            "validation_required_for": [
                "obligaciones",
                "verificaciones",
                "documentos",
                "facturas conformadas o pagadas",
                "penalidades",
                "modificaciones y prorrogas aplicadas",
                "expedientes",
            ],
            "feedback_enabled": True,
        },
        "security": {
            "auth_mode": config.auth_mode,
            "permissions": {role: sorted(perms) for role, perms in ROLE_PERMISSIONS.items()},
            "export_requires_permission": "export",
            "sensitive_actions_require_permission": "validate",
            "tenant_isolation": "tenant_id obligatorio en consultas y mutaciones operativas",
            "deployment_controls": {
                "ens_target": "Nivel bajo para seguimiento contractual; nivel medio cuando el despliegue trate datos personales o informacion confidencial.",
                "hosting_region": "El despliegue productivo debe ubicarse en la UE y acreditarse por el operador de infraestructura.",
                "closed_ai_environment": "El proveedor/modelo configurado debe garantizar no reutilizacion para entrenamiento, limitacion de retencion y confidencialidad.",
            },
        },
        "data_governance": {
            "postgres_schema": "xfollow",
            "data_protection_scan": "Escaneo local de evidencias y ficha contractual para detectar datos personales, IBAN, telefonos, correos y senales de confidencialidad.",
            "object_storage": {
                "endpoint": config.object_storage_endpoint,
                "bucket": config.object_storage_bucket,
                "local_fallback": "data/evidence/{tenant}/{contract}",
            },
            "secrets": "Las claves y DSN se redaccionan en respuestas de configuracion.",
            "training_reuse": "El diseno asume entorno cerrado o proveedor sin reutilizacion de datos; debe acreditarse en despliegue.",
            "retention": "La retencion y borrado se parametrizan por politica del organismo; el export permite devolucion completa.",
        },
        "ai_literacy": {
            "resources": ["docs/USER_GUIDE.md", "docs/COMPLIANCE_MATRIX.md", "README.md"],
            "operator_training": "Guia funcional, limites de IA, validacion humana, feedback, proteccion de datos y devolucion del servicio.",
        },
        "service_return": {
            "formats": ["json", "zip", "docx", "markdown"],
            "included": [
                "contrato",
                "obligaciones",
                "alertas",
                "evidencias",
                "verificaciones",
                "facturas",
                "liquidacion",
                "cambios",
                "documentos",
                "penalidades",
                "expedientes",
                "proteccion_datos",
                "feedback",
                "auditoria",
            ],
        },
    }


def list_audit_events(user: UserContext, contract_id: str) -> dict[str, Any]:
    if not get_contract(user, contract_id):
        raise KeyError(contract_id)
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.audit_events WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at DESC LIMIT 250",
            (user.tenant_id, contract_id),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["audit"].values() if item["tenant_id"] == user.tenant_id and item.get("contract_id") == contract_id]
        items.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("aud"), "items": items}


def list_tenant_audit_events(user: UserContext) -> dict[str, Any]:
    if db.db_available():
        rows = db.fetch_all(
            "SELECT * FROM xfollow.audit_events WHERE tenant_id = %s ORDER BY created_at DESC",
            (user.tenant_id,),
        )
        items = [_row(row) for row in rows]
    else:
        items = [item for item in _MEMORY["audit"].values() if item["tenant_id"] == user.tenant_id]
        items.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {"trace_id": trace_id("aud"), "items": items}


def contract_governance(user: UserContext, contract_id: str) -> dict[str, Any]:
    audit = list_audit_events(user, contract_id)["items"]
    feedback = list_feedback(user, contract_id)["items"]
    manifest = governance_manifest(user)
    actions: dict[str, int] = {}
    for event in audit:
        action = str(event.get("action") or "unknown")
        actions[action] = actions.get(action, 0) + 1
    feedback_summary: dict[str, int] = {}
    for item in feedback:
        rating = str(item.get("rating") or "neutral")
        feedback_summary[rating] = feedback_summary.get(rating, 0) + 1
    return {
        "trace_id": trace_id("gov"),
        "manifest": manifest,
        "audit_summary": {
            "events": len(audit),
            "actions": actions,
            "last_event_at": audit[0].get("created_at") if audit else None,
        },
        "feedback_summary": {
            "items": len(feedback),
            "ratings": feedback_summary,
        },
    }


def service_return_export(user: UserContext, *, audit_event: bool = True) -> dict[str, Any]:
    contracts = list_contracts(user, include_archived=True)["items"]
    if audit_event:
        _audit(user, "service_return.exported", None, {"contracts": len(contracts), "format": "json"})
    contract_exports = [export_contract(user, contract["id"]) for contract in contracts]
    tenant_audit = list_tenant_audit_events(user)["items"]
    manifest = governance_manifest(user)
    return {
        "trace_id": trace_id("ret"),
        "manifest": {
            "service": "xFollow",
            "version": "0.1.0",
            "tenant_id": user.tenant_id,
            "exported_at": now_iso(),
            "format": "service_return_json",
            "contracts": len(contracts),
        },
        "governance": manifest,
        "summary": {
            **_summary(user),
            "tenant_audit_events": len(tenant_audit),
            "generated_by": user.user_id,
            "included_formats": ["json", "zip", "docx", "markdown"],
        },
        "contracts": contracts,
        "contract_exports": contract_exports,
        "tenant_audit": tenant_audit,
        "transition_guide": _service_return_guide(len(contracts)),
        "transition_documents": _transition_documents(),
    }


def service_return_zip_bytes(user: UserContext) -> tuple[str, bytes]:
    export = service_return_export(user, audit_event=True)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", json.dumps(export["manifest"], ensure_ascii=False, indent=2))
        archive.writestr("governance.json", json.dumps(export["governance"], ensure_ascii=False, indent=2))
        archive.writestr("summary.json", json.dumps(export["summary"], ensure_ascii=False, indent=2))
        archive.writestr("contracts/index.json", json.dumps(export["contracts"], ensure_ascii=False, indent=2))
        archive.writestr("audit/tenant_audit.json", json.dumps(export["tenant_audit"], ensure_ascii=False, indent=2))
        archive.writestr("transition/README.md", export["transition_guide"])
        for filename, content in export["transition_documents"].items():
            archive.writestr(f"transition/{filename}", content)
        for contract_export in export["contract_exports"]:
            contract = contract_export["contract"]
            base = _safe_filename(contract.get("file_number") or contract["id"])
            archive.writestr(f"contracts/{base}/export.json", json.dumps(contract_export, ensure_ascii=False, indent=2))
            archive.writestr(f"contracts/{base}/contract.json", json.dumps(contract_export["contract"], ensure_ascii=False, indent=2))
            for key in [
                "obligations",
                "alerts",
                "evidence",
                "data_protection",
                "checks",
                "invoices",
                "liquidation",
                "changes",
                "penalties",
                "proceedings",
                "feedback",
                "comments",
                "governance",
            ]:
                archive.writestr(f"contracts/{base}/{key}.json", json.dumps(contract_export[key], ensure_ascii=False, indent=2))
            archive.writestr(f"contracts/{base}/audit.json", json.dumps(contract_export["audit"], ensure_ascii=False, indent=2))
            for document in contract_export["documents"]:
                doc_base = _safe_filename(f"{document['document_type']}-{document['id']}")
                text = document.get("content_text") or ""
                archive.writestr(f"contracts/{base}/documents/{doc_base}.md", text)
                archive.writestr(f"contracts/{base}/documents/{doc_base}.docx", _docx_bytes(document["title"], text))
    filename = f"{_safe_filename(user.tenant_id)}-xfollow-service-return.zip"
    return filename, buffer.getvalue()


def _service_return_guide(contract_count: int) -> str:
    return "\n".join(
        [
            "# Paquete de devolucion xFollow",
            "",
            "Este paquete contiene los datos generados y gestionados por xFollow para facilitar la continuidad o sustitucion del servicio.",
            "",
            "## Contenido",
            "",
            "- `manifest.json`: identificacion del paquete, tenant y fecha de exportacion.",
            "- `governance.json`: controles de IA, permisos, gobierno de datos y formatos.",
            "- `summary.json`: resumen operativo del tenant.",
            "- `contracts/index.json`: inventario de contratos exportados.",
            "- `contracts/*/export.json`: datos completos de cada contrato.",
            "- `contracts/*/data_protection.json`: evaluaciones de datos personales y confidencialidad.",
            "- `contracts/*/documents/*`: documentos generados en Markdown y DOCX.",
            "- `audit/tenant_audit.json`: eventos de auditoria del tenant.",
            "- `transition/USER_GUIDE.md`: guia funcional para continuidad del servicio.",
            "- `transition/COMPLIANCE_MATRIX.md`: matriz de cobertura funcional y controles transversales.",
            "",
            "## Reutilizacion",
            "",
            "Los JSON son UTF-8 y pueden cargarse en otra solucion o en herramientas analiticas. Los DOCX son copias revisables de los borradores generados.",
            "",
            "## Control",
            "",
            "La exportacion requiere permiso `export`. Los borradores generados por IA no tienen efecto administrativo sin validacion humana.",
            "",
            f"Contratos incluidos: {contract_count}.",
        ]
    )


def _transition_documents() -> dict[str, str]:
    docs_dir = repo_root() / "docs"
    documents: dict[str, str] = {}
    for filename in ["USER_GUIDE.md", "COMPLIANCE_MATRIX.md", "ARCHITECTURE.md"]:
        path = docs_dir / filename
        if path.exists():
            documents[filename] = path.read_text(encoding="utf-8", errors="replace")
    return documents


def export_contract(user: UserContext, contract_id: str) -> dict[str, Any]:
    contract = get_contract(user, contract_id)
    if not contract:
        raise KeyError(contract_id)
    if db.db_available():
        audit_rows = db.fetch_all("SELECT * FROM xfollow.audit_events WHERE tenant_id = %s AND contract_id = %s ORDER BY created_at", (user.tenant_id, contract_id))
        audit = [_row(row) for row in audit_rows]
    else:
        audit = [item for item in _MEMORY["audit"].values() if item["tenant_id"] == user.tenant_id and item.get("contract_id") == contract_id]
    return {
        "trace_id": trace_id("exp"),
        "manifest": {"service": "xFollow", "version": "0.1.0", "exported_at": now_iso(), "format": "json"},
        "contract": contract,
        "obligations": list_obligations(user, contract_id)["items"],
        "alerts": list_alerts(user, contract_id)["items"],
        "evidence": list_evidence(user, contract_id)["items"],
        "data_protection": list_data_protection_assessments(user, contract_id)["items"],
        "checks": list_checks(user, contract_id)["items"],
        "invoices": list_invoices(user, contract_id)["items"],
        "liquidation": liquidation_preview(user, contract_id)["summary"],
        "changes": list_contract_changes(user, contract_id)["items"],
        "documents": list_documents(user, contract_id)["items"],
        "penalties": list_penalties(user, contract_id)["items"],
        "proceedings": list_proceedings(user, contract_id)["items"],
        "feedback": list_feedback(user, contract_id)["items"],
        "comments": list_contract_comments(user, contract_id)["items"],
        "governance": contract_governance(user, contract_id),
        "audit": audit,
    }


def document_docx_bytes(user: UserContext, document_id: str) -> tuple[str, bytes]:
    document = _get_document(user, document_id)
    if not document:
        raise KeyError(document_id)
    filename = f"{_safe_filename(document['title'])}.docx"
    return filename, _docx_bytes(document["title"], document.get("content_text") or "")


def export_contract_zip_bytes(user: UserContext, contract_id: str) -> tuple[str, bytes]:
    export = export_contract(user, contract_id)
    contract = export["contract"]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", json.dumps(export["manifest"], ensure_ascii=False, indent=2))
        archive.writestr("contract.json", json.dumps(contract, ensure_ascii=False, indent=2))
        archive.writestr("obligations.json", json.dumps(export["obligations"], ensure_ascii=False, indent=2))
        archive.writestr("alerts.json", json.dumps(export["alerts"], ensure_ascii=False, indent=2))
        archive.writestr("evidence.json", json.dumps(export["evidence"], ensure_ascii=False, indent=2))
        archive.writestr("data_protection.json", json.dumps(export["data_protection"], ensure_ascii=False, indent=2))
        archive.writestr("checks.json", json.dumps(export["checks"], ensure_ascii=False, indent=2))
        archive.writestr("invoices.json", json.dumps(export["invoices"], ensure_ascii=False, indent=2))
        archive.writestr("liquidation.json", json.dumps(export["liquidation"], ensure_ascii=False, indent=2))
        archive.writestr("changes.json", json.dumps(export["changes"], ensure_ascii=False, indent=2))
        archive.writestr("penalties.json", json.dumps(export["penalties"], ensure_ascii=False, indent=2))
        archive.writestr("proceedings.json", json.dumps(export["proceedings"], ensure_ascii=False, indent=2))
        archive.writestr("feedback.json", json.dumps(export["feedback"], ensure_ascii=False, indent=2))
        archive.writestr("comments.json", json.dumps(export["comments"], ensure_ascii=False, indent=2))
        archive.writestr("governance.json", json.dumps(export["governance"], ensure_ascii=False, indent=2))
        archive.writestr("audit.json", json.dumps(export["audit"], ensure_ascii=False, indent=2))
        for item in export["documents"]:
            base = _safe_filename(f"{item['document_type']}-{item['id']}")
            text = item.get("content_text") or ""
            archive.writestr(f"documents/{base}.md", text)
            archive.writestr(f"documents/{base}.docx", _docx_bytes(item["title"], text))
    filename = f"{_safe_filename(contract.get('file_number') or contract['id'])}-xfollow.zip"
    return filename, buffer.getvalue()


def _docx_bytes(title: str, markdown_text: str) -> bytes:
    document = Document()
    document.add_heading(title, level=1)
    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
